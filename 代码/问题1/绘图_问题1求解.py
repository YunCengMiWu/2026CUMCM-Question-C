#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""2026 高教社杯 C 题《微网与外部电网电力调控策略》· 问题1 求解可视化集。

产出 6 张图（PNG 300 dpi + SVG）到
    04_结果与检验/可视化集/01_问题1_求解/
运行日志写入
    04_结果与检验/检验日志/可视化_问题1求解.txt

本脚本：
  · 用项目自带模型（03_代码/最小运行/模型.py，方案丙口径）重解问题1，保证逐槽量自洽；
  · 与 04_结果与检验/results/result1.xlsx 逐项对照；
  · 只读既有文件，不修改/覆盖任何既有文件；
  · 结尾打印 PASS/FAIL 自检结论，全部通过时退出码 0。
"""
from __future__ import annotations

import sys
import traceback
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(r"C:\Users\John\Desktop\26题")
CODE_DIR = ROOT / "03_代码" / "最小运行"
OUT_DIR = ROOT / "04_结果与检验" / "可视化集" / "01_问题1_求解"
LOG_DIR = ROOT / "04_结果与检验" / "检验日志"
LOG_PATH = LOG_DIR / "可视化_问题1求解.txt"
ATT1 = ROOT / "00_题目与原始数据" / "C题" / "附件" / "附件1.xlsx"
RESULT1 = ROOT / "04_结果与检验" / "results" / "result1.xlsx"
SKILL_SCRIPTS = r"C:\Users\John\.codex\skills\math-modeling\tools\figure\scripts"

# ---------------------------------------------------------------- 期望值（题面给定）
EXP_OBJECTIVE = 35126.84858928963          # 元
EXP_TOTAL_BUY = 59482.6989983539           # kWh
EXP_TOTAL_C = 20740.6661                   # kWh
EXP_TOTAL_D = 16799.9396                   # kWh
EXP_SEG_ORIG_C = [4166.6667, 833.3333, 4787.9643, 5286.0352, 0.0, 5666.6667]
EXP_SEG_ORIG_D = [0.0, 6365.8412, 1702.9970, 91.1014, 5780.1319, 2859.8681]
EXP_SEG_CLOCK_C = [4500.0000, 833.3333, 3954.6310, 6119.3685, 0.0, 5333.3333]
EXP_SEG_CLOCK_D = [0.0, 5947.4198, 2121.4184, 91.1014, 5068.6869, 3571.3131]
EXP_TABLE1 = [0.0, 486.4029, 0.0, 394.9315, 636.9826, 0.0]
EXP_TABLE1_SLOTS = [60, 72, 84, 96, 108, 120]
EXP_STRICT_OBJECTIVE = 35126.948589290     # 严格口径 E0=E143=E144=6000
EXP_B0 = 48052.05                          # 元（无储能基线）
EXP_B0_CURTAIL = 6247.963                  # kWh（B0 基线弃光，非最优解弃光）
EXP_B1 = 45719.85                          # 元（贪心谷充峰放规则基线，末态 10800）
EXP_B1_SOC_END = 10800.0
EXP_E0 = 6300.0
EXP_E143 = 6000.0
EXP_PRICE_RANGE = (0.3713, 1.3952)
EXP_LOAD_RANGE = (3309.4, 5959.0)
EXP_PV_RANGE = (0.0, 7612.3)

SEG_LABELS = ["0:00-4:00", "4:00-8:00", "8:00-12:00",
              "12:00-16:00", "16:00-20:00", "20:00-24:00"]
SEG_ORIG_IDX = [(0, 24), (24, 48), (48, 72), (72, 96), (96, 120), (120, 144)]
SEG_CLOCK_IDX = [[143] + list(range(0, 23)), list(range(23, 47)), list(range(47, 71)),
                 list(range(71, 95)), list(range(95, 119)), list(range(119, 143))]

CHECKS: list[tuple[str, bool, str]] = []
FIG_FILES: list[Path] = []


def record(name: str, ok: bool, detail: str = "") -> bool:
    CHECKS.append((name, bool(ok), detail))
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""))
    return bool(ok)


class Tee:
    """同时写 stdout 与日志文件。"""

    def __init__(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self._f = open(path, "w", encoding="utf-8")

    def write(self, s: str) -> int:
        sys.__stdout__.write(s)
        self._f.write(s)
        return len(s)

    def flush(self) -> None:
        sys.__stdout__.flush()
        self._f.flush()

    def close(self) -> None:
        self._f.flush()
        self._f.close()


# ---------------------------------------------------------------- 数据与模型
def read_attachment1():
    import numpy as np
    from openpyxl import load_workbook

    ws = load_workbook(ATT1, data_only=True).worksheets[0]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    rows = [r for r in rows if r[0] is not None]
    label_rows = [str(r[0])[:8] for r in rows]
    price = np.array([float(r[1]) for r in rows])
    load_kw = np.array([float(r[2]) for r in rows])
    pv_kw = np.array([float(r[3]) for r in rows])
    return label_rows, price, load_kw, pv_kw


def read_result1():
    from openpyxl import load_workbook

    wb = load_workbook(RESULT1, data_only=True)
    plan = [r[1] for r in wb["计划购电量"].iter_rows(min_row=2, values_only=True)]
    plan = [float(v) for v in plan if v is not None]
    cd = [r for r in wb["充放电量"].iter_rows(min_row=2, values_only=True)]
    return plan, cd


# ---------------------------------------------------------------- 坐标与刻度工具
def hhmm(minutes: float) -> str:
    m = int(round(minutes)) % 1440
    return f"{m // 60}:{m % 60:02d}"


def clock_start_minutes(t: int) -> int:
    """槽索引 t（0 基，对应 1 基槽号 t+1）在时钟轴上的起始分钟（周期化到 0:00—24:00）。"""
    return (10 * (t + 1)) % 1440


def clock_order():
    """按时钟顺序排列的槽索引：槽144（0:00—0:10）在最前，其后 槽1…槽143。"""
    return [143] + list(range(143))


def periodic(series):
    """把长度 144 的逐槽量按时钟顺序（x = 0,10,…,1430 min）重排。"""
    import numpy as np

    return np.asarray(series, dtype=float)[clock_order()]


def soc_curve(E_all):
    """周期化 SOC 曲线：0:00—24:00 共 145 点。"""
    import numpy as np

    E = np.asarray(E_all, dtype=float)
    return np.arange(145) * 10.0, np.concatenate([[E[143]], E[0:144]])


def set_time_axis(ax, step_hours: int = 2, label: str = "时刻"):
    ax.set_xlim(0, 1440)
    ticks = list(range(0, 1441, step_hours * 60))
    ax.set_xticks(ticks)
    ax.set_xticklabels([hhmm(t) for t in ticks], rotation=0)
    ax.set_xlabel(label)


def seg_sum(arr, idx):
    return float(sum(float(arr[i]) for i in idx))


def layout_audit(fig, tol_px: float = 6.0):
    """图内纯文字框的（重叠对，越界项）自检；不含标注箭头包围盒，避免误报。"""
    import matplotlib.text as mtext

    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    tick_ids = set()
    for ax in fig.axes:
        for tl in (*ax.get_xticklabels(), *ax.get_yticklabels(),
                   *ax.get_xticklabels(minor=True), *ax.get_yticklabels(minor=True)):
            tick_ids.add(id(tl))
    items = []
    for t in fig.findobj(mtext.Text):
        if id(t) in tick_ids:
            continue
        if not t.get_visible() or not t.get_text().strip():
            continue
        items.append((t.get_text().strip().replace("\n", " ")[:40],
                      mtext.Text.get_window_extent(t, renderer)))
    W, H = float(fig.bbox.width), float(fig.bbox.height)
    clipped = [nm for nm, bb in items
               if bb.x0 < -2 or bb.y0 < -2 or bb.x1 > W + 2 or bb.y1 > H + 2]
    overlaps_found = []
    for i in range(len(items)):
        for j in range(i + 1, len(items)):
            A, B = items[i][1], items[j][1]
            ox = min(A.x1, B.x1) - max(A.x0, B.x0)
            oy = min(A.y1, B.y1) - max(A.y0, B.y0)
            if ox > tol_px and oy > tol_px:
                overlaps_found.append((items[i][0], items[j][0]))
    return overlaps_found, clipped


# ---------------------------------------------------------------- 绘图
def fig_model_structure(q1, params: dict, colors):
    """图1-1 模型结构示意：三要素（决策变量 / 目标 / 约束）+ 能量流。"""
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Rectangle

    fig = plt.figure(figsize=(11.5, 7.9))
    fig.set_layout_engine("none")
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")

    def panel(x0, y0, x1, y1, title, face):
        ax.add_patch(Rectangle((x0, y0), x1 - x0, y1 - y0, transform=ax.transAxes,
                               facecolor=face, edgecolor="#9aa0a6", lw=1.0, zorder=0))
        ax.text((x0 + x1) / 2, y1 - 0.018, title, transform=ax.transAxes,
                ha="center", va="top", fontsize=10.5, fontweight="bold", zorder=5)
        return y1 - 0.052

    def node(cx, cy, w, h, text, face, edge, fs=9.0):
        ax.add_patch(FancyBboxPatch((cx - w / 2, cy - h / 2), w, h,
                                    boxstyle="round,pad=0.006,rounding_size=0.012",
                                    transform=ax.transAxes, facecolor=face,
                                    edgecolor=edge, lw=1.2, zorder=3))
        ax.text(cx, cy, text, transform=ax.transAxes, ha="center", va="center",
                fontsize=fs, zorder=4)

    def arrow(p, q, color, lw=1.5, style="-|>", ms=11, rad=0.0):
        ax.add_patch(FancyArrowPatch(p, q, transform=ax.transAxes, arrowstyle=style,
                                     mutation_scale=ms, lw=lw, color=color,
                                     shrinkA=1.0, shrinkB=1.0, zorder=2,
                                     connectionstyle=f"arc3,rad={rad}"))

    c_buy, c_ld, c_pv, c_dis, c_w, c_e = (colors[0], colors[1], colors[2],
                                          colors[3], colors[4], colors[8])

    # ============ (a) 能量流与决策变量 ============
    ax.text(0.5, 0.985, "图 1-1（a）　问题 1 模型的能量流与决策变量（单位：kWh / 元）",
            transform=ax.transAxes, ha="center", va="top", fontsize=11, fontweight="bold")

    bus_y = 0.828
    ax.plot([0.055, 0.945], [bus_y, bus_y], transform=ax.transAxes,
            color="#37474f", lw=3.0, solid_capstyle="butt", zorder=1)
    ax.text(0.058, bus_y + 0.012, "微网母线（逐槽电量平衡节点）", transform=ax.transAxes,
            ha="left", va="bottom", fontsize=8.5, color="#37474f")

    node(0.165, 0.930, 0.15, 0.062, "外部电网", "#e8f0fe", c_buy)
    node(0.415, 0.930, 0.15, 0.062, "光伏发电", "#e6f4ea", c_pv)
    node(0.845, 0.930, 0.17, 0.062, "小区负载", "#fef7e0", c_ld)
    node(0.285, 0.700, 0.19, 0.062, "弃光 $w_t$", "#f3e8fd", c_w)
    node(0.615, 0.700, 0.24, 0.062, "储能设备", "#fce8e6", c_dis)

    arrow((0.165, 0.898), (0.165, bus_y), c_buy)
    arrow((0.415, 0.898), (0.415, bus_y), c_pv)
    arrow((0.845, bus_y), (0.845, 0.898), c_ld, style="<|-|>")
    arrow((0.285, bus_y), (0.285, 0.732), c_w)
    arrow((0.585, bus_y), (0.585, 0.734), c_dis, style="-|>")
    arrow((0.650, 0.734), (0.650, bus_y), c_dis, style="-|>")

    ax.text(0.178, 0.862, "购电量 $a_t$", transform=ax.transAxes, ha="left", va="center",
            fontsize=9, color=c_buy)
    ax.text(0.428, 0.862, "可用光伏 $v_t=\\Delta t\\,PV_t$", transform=ax.transAxes,
            ha="left", va="center", fontsize=9, color=c_pv)
    ax.text(0.812, 0.862, "负荷 $\\ell_t=\\Delta t\\,L_t$", transform=ax.transAxes,
            ha="right", va="center", fontsize=9, color=c_ld)
    ax.text(0.300, 0.766, "削减光伏", transform=ax.transAxes, ha="left", va="center",
            fontsize=8.5, color=c_w)
    ax.text(0.576, 0.780, "充电 $C_t\\rightarrow$", transform=ax.transAxes, ha="right",
            va="center", fontsize=9, color=c_dis)
    ax.text(0.660, 0.780, "$\\rightarrow$放电 $D_t$", transform=ax.transAxes, ha="left",
            va="center", fontsize=9, color=c_dis)
    ax.text(0.5, 0.640, "逐槽平衡：$a_t+v_t+D_t=\\ell_t+C_t+w_t$　（不可向电网售电，$a_t\\geq 0$）",
            transform=ax.transAxes, ha="center", va="center", fontsize=10,
            bbox=dict(boxstyle="round,pad=0.35", facecolor="#ffffff",
                      edgecolor="#9aa0a6", lw=1.0))

    # ============ (b) 决策变量 ============
    x0, x1 = 0.015, 0.375
    ytop = panel(x0, 0.035, x1, 0.600, "（b）决策变量（逐槽 $t=1,\\dots,144$）", "#f7f9fc")
    items = [
        ("$a_t$", "计划购电量", "kWh", c_buy),
        ("$C_t$", "储能充电量", "kWh", c_dis),
        ("$D_t$", "储能放电量", "kWh", c_dis),
        ("$w_t$", "光伏弃光量", "kWh", c_w),
        ("$E_t$", "槽末储电量", "kWh", c_e),
    ]
    yy = ytop - 0.012
    for sym, nm, unit, col in items:
        ax.text(x0 + 0.022, yy, sym, transform=ax.transAxes, ha="left", va="top",
                fontsize=13, color=col, fontweight="bold")
        ax.text(x0 + 0.085, yy - 0.004, nm, transform=ax.transAxes, ha="left", va="top",
                fontsize=10)
        ax.text(x1 - 0.018, yy - 0.006, f"[{unit}]", transform=ax.transAxes, ha="right",
                va="top", fontsize=9, color="#5f6368")
        yy -= 0.055
    ax.text(x0 + 0.022, yy - 0.004,
            "另有 0-1 变量 $z_t$：$z_t=1$ 表示该槽充电（充放互斥）；\n"
            "窗初态 $E_0$（0:10 储电量）在问题 1 中为自由变量。\n\n"
            "单位换算：$1$ 槽 $=\\Delta t=\\frac{1}{6}\\,$h，\n"
            "槽电量(kWh) $=$ 功率(kW) $\\times\\frac{1}{6}$。",
            transform=ax.transAxes, ha="left", va="top", fontsize=9, color="#3c4043")

    # ============ (c) 目标函数 ============
    ytop = panel(0.395, 0.490, 0.985, 0.600, "（c）目标函数", "#f7f9fc")
    ax.text(0.415, ytop - 0.004,
            "$\\min\\;\\sum_{t=1}^{144} p_t\\,a_t$　　（元）",
            transform=ax.transAxes, ha="left", va="top", fontsize=15)
    ax.text(0.700, ytop - 0.014,
            "$p_t$：第 $t$ 槽电价（元/kWh）；$a_t$：第 $t$ 槽购电量（kWh）\n"
            "问题 1 无紧急购电、无违约调整，故目标即全天购电费。",
            transform=ax.transAxes, ha="left", va="top", fontsize=9, color="#3c4043")

    # ============ (d) 约束 ============
    ytop = panel(0.395, 0.035, 0.985, 0.470, "（d）约束条件", "#f7f9fc")
    cons = [
        ("供能平衡", "$a_t+v_t+D_t=\\ell_t+C_t+w_t$", "kWh"),
        ("SOC 递推", "$E_t=E_{t-1}+0.9\\,C_t-D_t/0.9$", "kWh"),
        ("储电量区间", "$1200\\leq E_t\\leq 10800$", "kWh"),
        ("充放电上限", "$0\\leq C_t,\\,D_t\\leq 5000\\times\\frac{1}{6}=833.33$", "kWh"),
        ("充放互斥", "$C_t\\leq M z_t,\\quad D_t\\leq M(1-z_t),\\quad M=833.33$", "kWh"),
        ("不售电 / 弃光", "$a_t\\geq 0,\\quad 0\\leq w_t\\leq v_t$", "kWh"),
        ("日循环", "$E_{144}=E_0$（次日 0:10 复现当日 0:10 储电量）", "kWh"),
        ("字面锚定", "$E_{143}=6000$（24:00 储电量）", "kWh"),
    ]
    yy = ytop - 0.010
    for nm, expr, unit in cons:
        ax.text(0.412, yy, "·", transform=ax.transAxes, ha="left", va="top",
                fontsize=11, color=colors[0])
        ax.text(0.427, yy, nm, transform=ax.transAxes, ha="left", va="top", fontsize=9.5)
        ax.text(0.548, yy, expr, transform=ax.transAxes, ha="left", va="top", fontsize=10)
        ax.text(0.975, yy, f"[{unit}]", transform=ax.transAxes, ha="right", va="top",
                fontsize=8.5, color="#5f6368")
        yy -= 0.043
    ax.text(0.412, yy + 0.006,
            "在役解（本图数据源）：$E_0=E_{144}=6300$ kWh，$E_{143}=6000$ kWh，"
            "$\\sum_t w_t=0$（无弃光）；\n严格口径 $E_0=E_{143}=E_{144}=6000$ kWh 亦可行，"
            "对应最优购电费 35 126.9486 元。",
            transform=ax.transAxes, ha="left", va="top", fontsize=8.4, color="#3c4043")

    return fig


def fig_inputs_and_plan(labels, price, load_kw, pv_kw, a, colors):
    """图1-2 典型日输入（电价/负载/光伏）与逐槽计划购电量。"""
    import matplotlib.pyplot as plt
    import numpy as np

    fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(11.5, 8.0),
                                   gridspec_kw=dict(height_ratios=[1.0, 1.0]))
    xs = np.arange(144) * 10.0
    p_ord = periodic(price)
    l_ord = periodic(load_kw)
    v_ord = periodic(pv_kw)
    a_ord = periodic(a)

    ax1.plot(xs, l_ord, color=colors[1], lw=1.6, label="小区负载 $L_t$（kW）")
    ax1.fill_between(xs, 0, v_ord, step="post", color=colors[2], alpha=0.30)
    ax1.plot(xs, v_ord, color=colors[2], lw=1.6, label="光伏预测功率 $PV_t$（kW）")
    sur = v_ord > l_ord
    if sur.any():
        ax1.fill_between(xs, l_ord, v_ord, where=sur, step="post",
                         color=colors[2], alpha=0.55, lw=0)
    ax1.set_ylabel("功率（kW）")
    ax1.set_ylim(0, max(l_ord.max(), v_ord.max()) * 1.22)

    ax1b = ax1.twinx()
    ax1b.spines["right"].set_visible(True)
    ax1b.step(xs, p_ord, where="post", color=colors[4], lw=1.5,
              label="电价 $p_t$（元/kWh）")
    ax1b.set_ylabel("电价（元/kWh）", color=colors[4])
    ax1b.tick_params(axis="y", colors=colors[4])
    ax1b.set_ylim(0, p_ord.max() * 1.35)

    h1, l1 = ax1.get_legend_handles_labels()
    h2, l2 = ax1b.get_legend_handles_labels()
    ax1.legend(h1 + h2, l1 + l2, loc="upper left", ncol=3, frameon=True,
               framealpha=0.92, fontsize=8.5)
    set_time_axis(ax1, 2, "")
    ax1.set_title("图 1-2（a）　典型日输入：电价、小区负载与光伏发电预测功率（附件1，144 槽）",
                  fontsize=11)
    # 光伏盈余区间标注
    idx = np.where(sur)[0]
    if idx.size:
        ax1.annotate(f"光伏 $>$ 负载：{hhmm(xs[idx[0]])}—{hhmm(xs[idx[-1]] + 10)}"
                     f"（{idx.size} 槽，盈余 {np.sum(np.maximum(v_ord - l_ord, 0)) * 1 / 6:.1f} kWh）",
                     xy=(xs[idx[len(idx) // 2]], (l_ord[idx[len(idx) // 2]] + v_ord[idx[len(idx) // 2]]) / 2),
                     xytext=(xs[idx[0]] - 20, ax1.get_ylim()[1] * 0.60),
                     fontsize=8.5, color="#1e6b34",
                     arrowprops=dict(arrowstyle="-|>", color="#1e6b34", lw=1.0))

    ax2.bar(xs, a_ord, width=9.6, align="edge", color=colors[0], alpha=0.85,
            label="计划购电量 $a_t$（kWh/槽）")
    for s, val in zip(EXP_TABLE1_SLOTS, EXP_TABLE1):
        pos = (10 * s) % 1440
        ax2.bar([pos], [max(val, 1.0)], width=9.6, align="edge", color=colors[3],
                alpha=0.95, zorder=3)
        ax2.annotate(f"{val:.4f}", xy=(pos + 5, val), xytext=(pos + 5, val + 95),
                     ha="center", va="bottom", fontsize=8.5, color=colors[3],
                     arrowprops=dict(arrowstyle="-", color=colors[3], lw=0.8))
    from matplotlib.patches import Patch

    ax2.set_ylabel("购电量（kWh/槽）")
    ax2.set_ylim(0, a_ord.max() * 1.30)
    set_time_axis(ax2, 2, "时刻")
    ax2.legend(handles=[Patch(facecolor=colors[0], alpha=0.85,
                              label="计划购电量 $a_t$（kWh/槽）"),
                        Patch(facecolor=colors[3], alpha=0.95, edgecolor=colors[3],
                              label="表1 指定时段 10:00 / 12:00 / 14:00 / 16:00 / 18:00 / 20:00")],
               loc="upper left", frameon=True, framealpha=0.92, fontsize=8.5)
    ax2.set_title("图 1-2（b）　逐槽计划购电量：全天 %.4f kWh，购电费 %.4f 元；表1 六时段合计 %.4f kWh"
                  % (a.sum(), float(np.sum(price * a)), sum(EXP_TABLE1)), fontsize=11)
    ax2.text(0.995, 0.93,
             "时间轴口径：模型窗口 0:10—24:10，图中已周期化到 0:00—24:00\n"
             "（槽144 覆盖 0:00—0:10，其余槽 $j$ 覆盖 $10j$—$10j{+}10$ 分钟）",
             transform=ax2.transAxes, ha="right", va="top", fontsize=8, color="#5f6368")
    return fig


def fig_storage_soc(a, C, D, E_all, price, colors):
    """图1-3 储能逐槽充放电与 SOC 轨迹。"""
    import matplotlib.pyplot as plt
    import numpy as np

    fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(11.5, 8.4),
                                   gridspec_kw=dict(height_ratios=[1.0, 1.20]))
    xs = np.arange(144) * 10.0
    C_ord, D_ord = periodic(C), periodic(D)

    ax1.bar(xs, C_ord, width=9.6, align="edge", color=colors[0], alpha=0.9,
            label="充电量 $C_t$（kWh/槽，向上）")
    ax1.bar(xs, -D_ord, width=9.6, align="edge", color=colors[3], alpha=0.9,
            label="放电量 $D_t$（kWh/槽，向下）")
    ax1.axhline(0, color="#37474f", lw=0.9)
    ax1.axhline(833.33, color=colors[0], lw=0.9, ls=":")
    ax1.axhline(-833.33, color=colors[3], lw=0.9, ls=":")
    ax1.text(6, 848, "单槽充放上限 833.33 kWh $=5000\\,$kW$\\times\\frac{1}{6}$h",
             fontsize=8.2, va="bottom", ha="left", color=colors[0])
    ax1.set_ylabel("充 / 放电量（kWh/槽）")
    ax1.set_ylim(-1030, 1030)
    set_time_axis(ax1, 2, "")
    ax1.legend(loc="lower left", ncol=1, frameon=True, framealpha=0.92, fontsize=8.5)
    ax1.set_title("图 1-3（a）　储能逐槽充放电计划：充电合计 %.4f kWh（45 槽），"
                  "放电合计 %.4f kWh（37 槽）" % (C.sum(), D.sum()), fontsize=11)

    tx, ty = soc_curve(E_all)
    ax2.axhspan(1200, 10800, color="#eef2f7", zorder=0)
    ax2.axhspan(0, 1200, color="#fdecea", zorder=0)
    ax2.axhspan(10800, 13200, color="#fdecea", zorder=0)
    ax2.axhline(1200, color="#9aa0a6", lw=0.9, ls="--")
    ax2.axhline(10800, color="#9aa0a6", lw=0.9, ls="--")
    ax2.axhline(6000, color=colors[7], lw=1.1, ls=":")

    def _runs(idx):
        out, s, p = [], idx[0], idx[0]
        for k in idx[1:]:
            if k == p + 1:
                p = k
            else:
                out.append((s, p))
                s = p = k
        out.append((s, p))
        return out

    top = np.where(np.abs(ty - 10800) < 1e-6)[0]
    bot = np.where(np.abs(ty - 1200) < 1e-6)[0]
    if top.size:
        runs = _runs(top)
        main = max(runs, key=lambda r: r[1] - r[0])
        ax2.axvspan(tx[main[0]], tx[main[1]] + 10, color=colors[3], alpha=0.10, zorder=1)
        desc = "、".join(f"{hhmm(tx[i])}—{hhmm(tx[j] + 10)}" if j > i else hhmm(tx[i])
                         for i, j in runs)
        ax2.text(520, 12250, f"SOC 顶格 10800 kWh 的时段：{desc}", fontsize=8.6,
                 color="#8c1d18", ha="left", va="bottom")
    if bot.size:
        runs = _runs(bot)
        desc = "、".join(f"{hhmm(tx[i])}—{hhmm(tx[j] + 10)}" if j > i else hhmm(tx[i])
                         for i, j in runs)
        ax2.annotate(f"SOC 触底 1200 kWh（下限）：{desc}",
                     xy=(tx[bot[len(bot) // 2]], 1200), xytext=(1000, 2750),
                     ha="center", va="top", fontsize=8.6, color="#8c1d18",
                     arrowprops=dict(arrowstyle="-|>", color="#8c1d18", lw=1.0))

    ax2.plot(tx, ty, color=colors[0], lw=1.8, marker="o", ms=2.0, zorder=5,
             label="储电量 $E$（kWh，145 点）")
    ax2.text(840, 6150, "6000 kWh 参考线（附录1 的 0:00 初值 = 24:00 锚定值）",
             ha="left", va="bottom", fontsize=8.2, color=colors[7])
    ax2.text(1435, 10700, "上限 10800", ha="right", va="bottom", fontsize=8,
             color="#5f6368")
    ax2.text(20, 1290, "下限 1200", ha="left", va="bottom", fontsize=8, color="#5f6368")

    marks = [(0, ty[0], "E(0:00)=6000", (120, 2950), "left"),
             (10, ty[1], "E(0:10)=6300", (120, 4100), "left"),
             (960, ty[96], "16:00：$E$ 顶格 10800 kWh", (1005, 11380), "center"),
             (1440, ty[144], "E(24:00)=6000", (1442, 7500), "right")]
    for xm, val, txt, xytext, ha in marks:
        ax2.plot([xm], [val], marker="o", ms=6, mfc="white", mec=colors[3], mew=1.6,
                 zorder=7, clip_on=False)
        ax2.annotate(f"{txt}\n（{val:.0f} kWh）", xy=(xm, val), xytext=xytext, ha=ha,
                     va="top", fontsize=8.6, color="#202124",
                     arrowprops=dict(arrowstyle="-|>", color=colors[3], lw=1.0))
    ax2.set_ylim(120, 13200)
    ax2.set_ylabel("储电量 $E$（kWh）")
    set_time_axis(ax2, 2, "时刻")
    ax2.legend(loc="upper right", frameon=True, framealpha=0.92, fontsize=8.5)
    ax2.set_title("图 1-3（b）　储电量轨迹（周期化 0:00—24:00）：日循环 $E_{144}=E_0=6300$ kWh，"
                  "字面锚定 $E_{143}=6000$ kWh", fontsize=11)
    return fig


def fig_two_conventions(Co, Do, Cc, Dc, colors):
    """图1-4 表2 两种口径对照（原报口径 / 时钟口径）。"""
    import matplotlib.pyplot as plt
    import numpy as np

    fig = plt.figure(figsize=(11.5, 7.2))
    fig.set_layout_engine("none")
    axc = fig.add_axes([0.055, 0.360, 0.420, 0.520])
    axd = fig.add_axes([0.545, 0.360, 0.420, 0.520])
    axt = fig.add_axes([0.020, 0.010, 0.960, 0.300])
    axt.axis("off")

    x = np.arange(6)
    w = 0.40
    for ax, orig, clk, name, unit in ((axc, Co, Cc, "充电量", "kWh"),
                                      (axd, Do, Dc, "放电量", "kWh")):
        ymax = max(max(orig), max(clk)) * 1.40
        ax.set_ylim(0, ymax)
        b1 = ax.bar(x - w / 2, orig, w, color=colors[0], alpha=0.9, label="原报口径")
        b2 = ax.bar(x + w / 2, clk, w, color=colors[2], alpha=0.9, label="时钟口径")
        for i in range(6):
            changed = abs(clk[i] - orig[i]) > 5e-4
            ax.text(x[i] - w / 2, orig[i] + ymax * 0.010, f"{orig[i]:.2f}",
                    ha="center", va="bottom", fontsize=7.2, color=colors[0])
            ax.text(x[i] + w / 2, clk[i] + ymax * 0.010, f"{clk[i]:.2f}",
                    ha="center", va="bottom", fontsize=7.2,
                    color=colors[3] if changed else colors[2],
                    fontweight="bold" if changed else "normal")
            if changed:
                b2[i].set_edgecolor(colors[3])
                b2[i].set_linewidth(1.8)
                ax.text(x[i] + w / 2, max(orig[i], clk[i]) + ymax * 0.085,
                        "Δ %+.4f" % (clk[i] - orig[i]),
                        ha="center", va="bottom", fontsize=8.2, color=colors[3],
                        fontweight="bold")
        ax.set_xticks(x)
        ax.set_xticklabels(SEG_LABELS, fontsize=8.2)
        ax.set_ylabel(f"{name}（{unit}）")
        ax.legend(loc="upper center", ncol=2, fontsize=8.8, frameon=True,
                  framealpha=0.95, columnspacing=1.2)
        ax.set_title(f"（{'a' if name == '充电量' else 'b'}）{name}：两种口径对照",
                     fontsize=10.5)

    txt = (
        "口径差异说明：\n"
        "① 「原报口径」＝建模窗口的自然 24 槽分组（槽 1—24 / 25—48 / …）。窗口为 0:10—24:10，故「槽 1—24」实际覆盖 0:10—4:10，\n"
        "     与表2 的「0:00-4:00」相差 10 min；「槽 121—144」实际覆盖 20:10—24:10。\n"
        "② 「时钟口径」按真实时钟边界归集：0:00-4:00＝槽{144,1…23}、4:00-8:00＝槽{24…47}、8:00-12:00＝槽{48…71}、\n"
        "     12:00-16:00＝槽{72…95}、16:00-20:00＝槽{96…119}、20:00-24:00＝槽{120…143}，使表2 的时段标签字面成立。\n"
        "③ 两种口径全天合计完全相同：充电 20 740.6661 kWh、放电 16 799.9396 kWh；仅 8 个分项数值变化（红框 / Δ 标注），\n"
        "     差异全部来自 6 个跨边界槽：槽 24、48、72、96、120、144。\n"
        "④ 逐槽解、全天购电量 59 482.6990 kWh、购电费 35 126.8486 元、SOC 轨迹与弃光 0 均不受口径选择影响。\n"
        "⑤ **交付口径（2026-09-11 订正后）**：result1.xlsx 的「充放电量」表已改用「时钟口径」；"
        "原报口径仅作对照保留。"
    )
    axt.text(0.005, 0.98, txt, ha="left", va="top", fontsize=8.6, color="#202124",
             linespacing=1.5,
             bbox=dict(boxstyle="round,pad=0.5", facecolor="#f7f9fc",
                       edgecolor="#9aa0a6", lw=1.0))
    return fig


def fig_baselines(cost0, cost1, cost2, cur0, colors):
    """图1-5 三级基线阶梯。"""
    import matplotlib.pyplot as plt
    import numpy as np

    fig = plt.figure(figsize=(9.8, 7.0))
    fig.set_layout_engine("none")
    ax = fig.add_axes([0.105, 0.400, 0.875, 0.520])
    axt = fig.add_axes([0.020, 0.010, 0.960, 0.330])
    axt.axis("off")

    names = ["B0 无储能\n（纯购电）",
             "B1 规则基线\n（贪心谷充峰放）",
             "B2 本模型\n（问题1 MILP 最优）"]
    vals = [cost0, cost1, cost2]
    cols = [colors[7], colors[1], colors[2]]
    bars = ax.bar(names, vals, color=cols, alpha=0.92, width=0.52)
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width() / 2, v + max(vals) * 0.015, f"{v:,.2f} 元",
                ha="center", va="bottom", fontsize=10.5, fontweight="bold")
    ax.set_xlim(-0.62, 4.10)
    ax.set_ylim(0, max(vals) * 1.14)
    for v, c, lb in ((cost0, cols[0], "B0 水平"), (cost1, cols[1], "B1 水平")):
        ax.hlines(v, -0.45, 3.95, color=c, lw=1.1, ls=(0, (5, 3)))
        ax.text(1.50, v, lb, ha="center", va="center", fontsize=7.8, color=c,
                bbox=dict(boxstyle="round,pad=0.18", facecolor="white",
                          edgecolor="none", alpha=0.95))
    for xa, v0, lab, col in ((3.00, cost1, "较 B1 节省 10 593.00 元（−23.17%）", cols[1]),
                             (3.42, cost0, "较 B0 节省 12 925.20 元（−26.90%）", cols[0])):
        ax.annotate("", xy=(xa, v0), xytext=(xa, cost2),
                    arrowprops=dict(arrowstyle="<|-|>", color=col, lw=1.2))
        ax.text(xa + 0.055, (v0 + cost2) / 2, lab, rotation=90,
                ha="left", va="center", fontsize=8.8, color=col)
    ax.set_ylabel("全天购电费（元）")
    ax.grid(axis="y", ls=":", lw=0.6, alpha=0.6)
    ax.set_axisbelow(True)
    ax.set_title("图 1-5　问题 1 的三级基线阶梯（典型日，附件1 数据）", fontsize=11)

    txt = (
        "基线口径说明：\n"
        "· B0：不使用储能，逐槽按缺口购电，全天购电费 48 052.05 元；此时光伏盈余 6 247.963 kWh 全部弃掉——\n"
        "     6 247.963 kWh 属于 B0 基线，不是本模型解的弃光量。\n"
        "· B1：贪心「谷充峰放」规则基线，全天购电费 45 719.85 元，其末态储电量 10 800 kWh（上限顶格），\n"
        "     不满足问题1「0:00 与 24:00 储电量相同」的日循环要求，与 B2 不在同一可行域，仅作「未满足终端要求的规则示例」。\n"
        "· B2：本模型（方案丙口径 MILP）最优解，全天购电费 35 126.8486 元，弃光 0 kWh，\n"
        "     $E_{143}=6000$ kWh、$E_{144}=E_0=6300$ kWh，日循环成立。\n"
        "· 因此 B0 / B1 与 B2 的差额只能作为费用对照，不构成 B2 最优性的独立证明。"
    )
    axt.text(0.005, 0.98, txt, ha="left", va="top", fontsize=8.7, color="#202124",
             linespacing=1.5,
             bbox=dict(boxstyle="round,pad=0.5", facecolor="#f7f9fc",
                       edgecolor="#9aa0a6", lw=1.0))
    return fig


def fig_netload(load_kw, pv_kw, a, C, D, price, colors):
    """图1-6 净负荷与储能作用。"""
    import matplotlib.pyplot as plt
    import numpy as np

    DT = 1.0 / 6.0
    fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(11.5, 7.6), sharex=True,
                                   gridspec_kw=dict(height_ratios=[1.0, 0.85]))
    xs = np.arange(144) * 10.0
    ell = periodic(load_kw) * DT
    v = periodic(pv_kw) * DT
    net = ell - v
    a_o, C_o, D_o = periodic(a), periodic(C), periodic(D)

    for ax in (ax1, ax2):
        ax.axvspan(0, 360, color="#eceff1", alpha=0.55, zorder=0)
        ax.axvspan(1080, 1440, color="#eceff1", alpha=0.55, zorder=0)
    ax1.axhline(0, color="#37474f", lw=1.0)
    ax1.fill_between(xs, 0, net, where=(net >= 0), step="post", color=colors[1],
                     alpha=0.35, lw=0)
    ax1.fill_between(xs, 0, net, where=(net < 0), step="post", color=colors[2],
                     alpha=0.55, lw=0)
    ax1.step(xs, net, where="post", color="#202124", lw=1.3,
             label="净负荷 $\\ell_t-v_t$（kWh/槽）")
    ylo = float(net.min()) * 1.45
    yhi = float(net.max()) * 1.30
    ax1.set_ylim(ylo, yhi)
    ax1.set_ylabel("净负荷（kWh/槽）")
    ax1.set_title("图 1-6（a）　净负荷曲线：缺口段（购电 / 放电）与盈余段（充电 / 弃光）", fontsize=11)
    ax1.legend(loc="upper right", fontsize=8.5, frameon=True, framealpha=0.92)
    i_neg = np.where(net < 0)[0]
    tot_neg = float(-net[net < 0].sum())
    tot_pos = float(net[net > 0].sum())
    ax1.annotate(f"光伏盈余 {i_neg.size} 槽，合计 {tot_neg:,.2f} kWh\n"
                 f"（{hhmm(xs[i_neg[0]])}—{hhmm(xs[i_neg[-1]] + 10)}，全部被充电/负荷吸收，弃光 0）",
                 xy=(xs[i_neg[len(i_neg) // 2]], net[i_neg[len(i_neg) // 2]] * 0.55),
                 xytext=(xs[i_neg[0]] - 70, ylo * 0.72), fontsize=8.8, color="#1e6b34",
                 ha="left", arrowprops=dict(arrowstyle="-|>", color="#1e6b34", lw=1.0))
    ax1.text(20, yhi * 0.93,
             f"缺口 {int((net > 0).sum())} 槽，合计 {tot_pos:,.2f} kWh（由购电与放电补足）",
             fontsize=8.8, color="#8a4b00", ha="left", va="top")
    for xc, s in ((180, "夜间"), (720, "白天"), (1260, "夜间")):
        ax1.text(xc, yhi * 0.55, s, fontsize=9, color="#5f6368", ha="center",
                 bbox=dict(boxstyle="round,pad=0.25", facecolor="white",
                           edgecolor="none", alpha=0.75))

    ax2.bar(xs, a_o, width=9.6, align="edge", color=colors[0], alpha=0.75,
            label="购电量 $a_t$（kWh/槽）")
    ax2.bar(xs, D_o, width=9.6, align="edge", color=colors[3], alpha=0.85,
            label="放电量 $D_t$（kWh/槽）")
    ax2.bar(xs, -C_o, width=9.6, align="edge", color=colors[8], alpha=0.85,
            label="充电量 $C_t$（kWh/槽，向下）")
    ax2.axhline(0, color="#37474f", lw=1.0)
    ax2.set_ylabel("电量（kWh/槽）")
    ax2.set_ylim(-1050, 1900)
    set_time_axis(ax2, 2, "时刻")
    ax2.legend(loc="upper left", ncol=3, fontsize=8.4, frameon=True, framealpha=0.95,
               columnspacing=1.0, handlelength=1.4)
    ax2.set_title("图 1-6（b）　逐槽购电 / 放电（向上）与充电（向下）：盈余段由充电吸收，"
                  "缺口段由购电与放电补足", fontsize=11)
    for xm, ytxt, s, col in [(605, 900, "10:00 购电 0：光伏 1068.27 kWh > 负载 983.70 kWh", colors[2]),
                             (845, 1330, "14:00 购电 0：光伏 1068.75 kWh > 负载 948.72 kWh", colors[2]),
                             (1215, 1760, "20:00 购电 0：光伏为 0，由储能放电 711.45 kWh 覆盖", colors[3])]:
        ax2.annotate(s, xy=(xm, 90), xytext=(xm, ytxt), ha="center", va="bottom",
                     fontsize=8.2, color=col,
                     arrowprops=dict(arrowstyle="-|>", color=col, lw=0.9))
    ax2.text(0.995, 0.02,
             "阴影：夜间（0:00—6:00、18:00—24:00，作图约定）\n"
             "槽电量 = 功率 × (1/6) h；净负荷 $\\ell_t-v_t$ 为槽电量口径（kWh）",
             transform=ax2.transAxes, ha="right", va="bottom", fontsize=8, color="#5f6368")
    return fig


# ---------------------------------------------------------------- 主流程
def main() -> int:
    import numpy as np

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    LOG_DIR.mkdir(parents=True, exist_ok=True)

    print("=" * 100)
    print("2026 高教社杯 C 题 · 问题1 求解可视化集生成")
    print("=" * 100)
    print(f"工作目录      : {ROOT}")
    print(f"输出目录      : {OUT_DIR}")
    print(f"日志文件      : {LOG_PATH}")
    print(f"模型代码      : {CODE_DIR / '模型.py'}（方案丙口径，只读）")
    print(f"对照结果      : {RESULT1}（只读）")
    print("-" * 100)

    # ---- 样式 ----
    import matplotlib

    matplotlib.use("Agg")
    sys.path.insert(0, SKILL_SCRIPTS)
    style_backend = "手工 rcParams"
    try:
        from setup_style import setup_style

        info = setup_style(journal="general", lang="zh")
        style_backend = f"setup_style({info})"
    except Exception as exc:  # noqa: BLE001
        import matplotlib.pyplot as plt

        plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
        plt.rcParams["axes.unicode_minus"] = False
        style_backend = f"手工 rcParams（Skill 导入失败：{exc}）"
    print(f"绘图样式      : {style_backend}")

    import matplotlib.pyplot as plt

    try:
        from export_figure import export_figure
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"export_figure 不可用：{exc}") from exc

    colors = plt.get_cmap("tab10").colors

    # ---- 数据 ----
    labels, price, load_kw, pv_kw = read_attachment1()
    record("附件1 结构核对（144 数据行 × 4 列）",
           len(price) == 144 and len(load_kw) == 144 and len(pv_kw) == 144,
           f"行数={len(price)}，首行标签={labels[0]!r}，末行标签={labels[-1]!r}")

    # ---- 模型求解 ----
    sys.path.insert(0, str(CODE_DIR))
    from 模型 import solve_day

    q1 = solve_day(price, load_kw, pv_kw, E_start=None, mode="q1", E_clock_24=6000.0)
    a, C, D, w, E_all = q1["a"], q1["C"], q1["D"], q1["w"], q1["E_all"]
    obj = float(q1["objective"])
    print("-" * 100)
    print(f"模型重解：status={q1['status']} ok={q1['ok']} mip_gap={q1['mip_gap']} "
          f"wall={q1['wall_s']:.3f}s")
    print(f"  目标值 = {obj!r} 元；Σa = {a.sum():.6f} kWh；ΣC = {C.sum():.6f} kWh；"
          f"ΣD = {D.sum():.6f} kWh；Σw = {w.sum():.10f} kWh")
    print(f"  E0 = {E_all[0]:.4f}；E143（24:00）= {E_all[143]:.4f}；E144（24:10）= {E_all[144]:.4f}；"
          f"E 最小 {E_all.min():.4f} / 最大 {E_all.max():.4f}")

    # ---- 自检 1：目标值 ----
    rel = abs(obj - EXP_OBJECTIVE) / abs(EXP_OBJECTIVE)
    record("自检1 目标值核对（相对差 < 1e-6）", rel < 1e-6,
           f"重解 {obj!r} vs 题面 {EXP_OBJECTIVE!r}，相对差 {rel:.3e}")
    record("自检1b 全天购电量核对", abs(a.sum() - EXP_TOTAL_BUY) < 1e-3,
           f"{a.sum():.6f} kWh vs {EXP_TOTAL_BUY} kWh")
    record("自检1c 全天购电费 = Σ p·a", abs(float(np.sum(price * a)) - obj) < 1e-6,
           f"Σ p·a = {float(np.sum(price * a)):.9f} 元")
    record("自检1d 充/放电合计", abs(C.sum() - EXP_TOTAL_C) < 5e-4
           and abs(D.sum() - EXP_TOTAL_D) < 5e-4,
           f"ΣC={C.sum():.4f}（期望 {EXP_TOTAL_C}）ΣD={D.sum():.4f}（期望 {EXP_TOTAL_D}）")
    record("自检1e 弃光 = 0（本最优解）", abs(w.sum()) < 1e-9,
           f"Σw = {w.sum():.10f} kWh")

    # ---- 自检 2：六段汇总 vs result1.xlsx（现行文件已按真实时钟区间订正）----
    plan_file, cd_file = read_result1()
    Co = [seg_sum(C, range(i, j)) for i, j in SEG_ORIG_IDX]      # 原报口径（历史对照）
    Do = [seg_sum(D, range(i, j)) for i, j in SEG_ORIG_IDX]
    Cc = [seg_sum(C, idx) for idx in SEG_CLOCK_IDX]              # 真实时钟区间（题目表2）
    Dc = [seg_sum(D, idx) for idx in SEG_CLOCK_IDX]
    file_co = [float(r[1]) for r in cd_file[:6]]
    file_do = [float(r[2]) for r in cd_file[:6]]
    file_lab = [str(r[0]) for r in cd_file[:6]]
    d_file = max(max(abs(x - y) for x, y in zip(Cc, file_co)),
                 max(abs(x - y) for x, y in zip(Dc, file_do)))
    d_orig = max(max(abs(x - y) for x, y in zip(Co, file_co)),
                 max(abs(x - y) for x, y in zip(Do, file_do)))
    print("  [文件段标签] :", file_lab)
    for k in range(6):
        print(f"    {SEG_LABELS[k]:<12} 时钟 {Cc[k]:10.4f}（文件 {file_co[k]:10.4f}）"
              f"  放 {Dc[k]:10.4f}（文件 {file_do[k]:10.4f}）  Δ充 {Cc[k]-Co[k]:+9.4f}")
    record("自检2 result1.xlsx「充放电量」表 = 真实时钟区间汇总（差 < 5e-4）",
           d_file < 5e-4, f"最大偏差 {d_file:.3e}")
    record("自检2b 该表已不再是原报口径（24 槽自然分组）",
           d_orig > 1e-2, f"与原报口径最大偏差 {d_orig:.4f} kWh")
    record("自检2c 六段标签与题目表2 时钟区间逐字一致",
           file_lab == SEG_LABELS, f"{file_lab}")
    d_exp = max(max(abs(x - y) for x, y in zip(Cc, EXP_SEG_CLOCK_C)),
                max(abs(x - y) for x, y in zip(Dc, EXP_SEG_CLOCK_D)))
    record("自检2d 时钟口径六段 = 题面给定值", d_exp < 5e-4, f"最大偏差 {d_exp:.3e}")

    # ---- 自检 3：计划购电量 144 值与 result1.xlsx 一致 ----
    plan_file = np.array(plan_file, dtype=float)[:144]
    dp = np.abs(a - plan_file)
    n_bad = int((dp > 5e-4).sum())
    record("自检3 「计划购电量」144 值逐槽一致（容差 5e-4，对应 4 位小数存储舍入）",
           n_bad == 0, f"最大偏差 {dp.max():.3e}，超容差槽数 {n_bad}")

    # ---- 自检 4：两种口径的差值与合计守恒 ----
    for k in range(6):
        print(f"    [对照] {SEG_LABELS[k]:<12} 原报 充 {Co[k]:10.4f} 放 {Do[k]:10.4f}"
              f"  →  时钟 充 {Cc[k]:10.4f} 放 {Dc[k]:10.4f}")
    n_changed = sum(1 for k in range(6) if abs(Cc[k] - Co[k]) > 5e-4) + \
        sum(1 for k in range(6) if abs(Dc[k] - Do[k]) > 5e-4)
    record("自检4 两口径全天合计一致（仅归组不同）", abs(sum(Cc) - sum(Co)) < 1e-6
           and abs(sum(Dc) - sum(Do)) < 1e-6,
           f"Σ充 {sum(Cc):.4f} / {sum(Co):.4f}；Σ放 {sum(Dc):.4f} / {sum(Do):.4f}")
    record("自检4b 12 个汇总数中变化的分项数 = 8", n_changed == 8, f"实测变化分项 {n_changed} 个")
    soc_c = EXP_E143
    for k in range(6):
        soc_c = soc_c + 0.9 * Cc[k] - Dc[k] / 0.9
    record("自检4c 时钟口径 SOC 链自 0:00=6000 闭环到 24:00=6000", abs(soc_c - EXP_E143) < 1e-6,
           f"链末端 {soc_c:.6f} kWh")

    # ---- 自检 5：SOC 检查点 ----
    record("自检5 SOC 检查点 E(0:10)=E_0=6300、E(24:00)=E_143=6000、E_144=E_0",
           abs(E_all[0] - EXP_E0) < 1e-6 and abs(E_all[143] - EXP_E143) < 1e-6
           and abs(E_all[144] - E_all[0]) < 1e-6,
           f"E0={E_all[0]:.4f} E143={E_all[143]:.4f} E144={E_all[144]:.4f}")
    record("自检5b E ∈ [1200, 10800]",
           E_all.min() >= 1200 - 1e-6 and E_all.max() <= 10800 + 1e-6,
           f"[{E_all.min():.4f}, {E_all.max():.4f}]")

    # ---- 自检 6：表1 六时段 ----
    a6 = [float(a[s - 1]) for s in EXP_TABLE1_SLOTS]
    d6 = max(abs(x - y) for x, y in zip(a6, EXP_TABLE1))
    record("自检6 表1 六时段（槽 60/72/84/96/108/120）购电量", d6 < 5e-4,
           " / ".join(f"{v:.4f}" for v in a6) + f"（最大偏差 {d6:.3e}）")
    for s, v in zip(EXP_TABLE1_SLOTS, a6):
        t = s - 1
        print(f"    槽{s:>4} {hhmm(10 * s)}—{hhmm(10 * s + 10)}  a={v:9.4f} kWh"
              f"  C={C[t]:8.4f}  D={D[t]:8.4f}"
              f"  光伏 {pv_kw[t]:8.2f} kW / {pv_kw[t] / 6:8.3f} kWh"
              f"  负载 {load_kw[t]:8.2f} kW / {load_kw[t] / 6:8.3f} kWh"
              f"  电价 {price[t]:.4f}")

    # ---- 自检 7：严格口径 ----
    q1s = solve_day(price, load_kw, pv_kw, E_start=6000.0, mode="q1", E_clock_24=6000.0)
    rel_s = abs(float(q1s["objective"]) - EXP_STRICT_OBJECTIVE) / EXP_STRICT_OBJECTIVE
    record("自检7 严格口径（E0=E143=E144=6000）最优值", rel_s < 1e-6,
           f"{q1s['objective']:.9f} 元 vs {EXP_STRICT_OBJECTIVE}（相对差 {rel_s:.3e}）")

    # ---- 自检 8：基线 ----
    DT = 1.0 / 6.0
    net_kwh = load_kw * DT - pv_kw * DT
    b0 = float(np.sum(price * np.maximum(net_kwh, 0.0)))
    cur0 = float(np.maximum(-net_kwh, 0.0).sum())
    record("自检8 B0 基线（无储能纯购电 + 弃光）",
           abs(b0 - EXP_B0) < 5e-3 and abs(cur0 - EXP_B0_CURTAIL) < 5e-3,
           f"购电费 {b0:.6f} 元（≈{b0:.2f}），弃光 {cur0:.6f} kWh（≈{cur0:.3f}，归属 B0）")

    # ---- 自检 9：输入数据范围 ----
    rng_ok = (abs(price.min() - EXP_PRICE_RANGE[0]) < 5e-4
              and abs(price.max() - EXP_PRICE_RANGE[1]) < 5e-4
              and abs(load_kw.min() - EXP_LOAD_RANGE[0]) < 5e-2
              and abs(load_kw.max() - EXP_LOAD_RANGE[1]) < 5e-2
              and abs(pv_kw.min() - EXP_PV_RANGE[0]) < 5e-2
              and abs(pv_kw.max() - EXP_PV_RANGE[1]) < 5e-2)
    record("自检9 输入数据范围（电价/负载/光伏）", rng_ok,
           f"电价 {price.min():.4f}—{price.max():.4f} 元/kWh；"
           f"负载 {load_kw.min():.1f}—{load_kw.max():.1f} kW；"
           f"光伏 {pv_kw.min():.1f}—{pv_kw.max():.1f} kW")

    # ---- 绘图 ----
    print("-" * 100)
    print("开始绘图（PNG 300 dpi + SVG）……")
    figs = []
    figs.append(("图1-1_模型结构示意", fig_model_structure(q1, {}, colors)))
    figs.append(("图1-2_典型日输入与购电计划",
                 fig_inputs_and_plan(labels, price, load_kw, pv_kw, a, colors)))
    figs.append(("图1-3_储能调度与SOC轨迹",
                 fig_storage_soc(a, C, D, E_all, price, colors)))
    figs.append(("图1-4_表2两种口径对照", fig_two_conventions(Co, Do, Cc, Dc, colors)))
    figs.append(("图1-5_基线阶梯", fig_baselines(b0, EXP_B1, obj, cur0, colors)))
    figs.append(("图1-6_净负荷与储能作用",
                 fig_netload(load_kw, pv_kw, a, C, D, price, colors)))

    print("-" * 100)
    layout_ok = True
    for name, fig in figs:
        base = OUT_DIR / name
        ov, cl = layout_audit(fig)
        ok = not ov and not cl
        layout_ok &= ok
        print(f"  [{'PASS' if ok else 'FAIL'}] {name} 版面自检"
              + (f" — 文字重叠 {len(ov)} 对，越界 {len(cl)} 项" if not ok
                 else " — 文字重叠 0 对，越界 0 项"))
        if not ok:
            for o in ov[:5]:
                print(f"        重叠：{o[0]!r} <-> {o[1]!r}")
            for c in cl[:5]:
                print(f"        越界：{c!r}")
        export_figure(fig, str(base), formats=("png", "svg"), dpi=300, tight=True)
        plt.close(fig)
        FIG_FILES.extend([Path(f"{base}.png"), Path(f"{base}.svg")])
    record("自检10 六张图版面自检（纯文字框 0 重叠、0 越界）", layout_ok,
           f"共检查 {len(figs)} 张图")

    # ---- 自检 10：图件字节数 ----
    print("-" * 100)
    print("图件字节数核对（要求 > 20000 字节）：")
    size_ok = True
    for name, _ in figs:
        png = OUT_DIR / f"{name}.png"
        svg = OUT_DIR / f"{name}.svg"
        pb = png.stat().st_size if png.exists() else 0
        sb = svg.stat().st_size if svg.exists() else 0
        ok = pb > 20000 and sb > 20000
        size_ok &= ok
        print(f"  [{'PASS' if ok else 'FAIL'}] {name}.png {pb:>9,d} 字节  |  "
              f"{name}.svg {sb:>9,d} 字节")
    record("自检11 每张图 PNG/SVG 字节数 > 20000", size_ok,
           f"{len(figs)} 张图，共 {len(FIG_FILES)} 个文件")

    # ---- 汇总 ----
    print("=" * 100)
    n_fail = sum(1 for _, ok, _ in CHECKS if not ok)
    for name, ok, detail in CHECKS:
        print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""))
    print("-" * 100)
    print(f"自检合计：{len(CHECKS) - n_fail}/{len(CHECKS)} PASS，{n_fail} FAIL")
    print(f"PNG/SVG 文件共 {len(FIG_FILES)} 个，位于 {OUT_DIR}")
    print("=" * 100)
    if n_fail:
        print("最终结论：FAIL")
        return 1
    print("最终结论：PASS（全部自检通过，退出码 0）")
    return 0


def run() -> int:
    tee = Tee(LOG_PATH)
    old = sys.stdout
    sys.stdout = tee
    try:
        return main()
    except Exception:  # noqa: BLE001
        traceback.print_exc(file=sys.stdout)
        print("=" * 100)
        print("最终结论：FAIL（脚本异常终止）")
        return 1
    finally:
        sys.stdout = old
        tee.close()
        print(f"运行日志已写入：{LOG_PATH}")


if __name__ == "__main__":
    sys.exit(run())
