#!/usr/bin/env python3
"""F1 正式模板落盘：按官方 5 个模板（result1/result2/result3/result4-2/result4-3）填写全年结果。

口径（与 §13/§14 及方案丙一致，供论文"结果说明"引用）：
- 计划购电量 / 调整购电量：144 槽（列标签沿用模板：0:10-0:20 … 23:50-0:00+1、0:00-0:10+1），
  末两列为"全天购电量"（Σ 该行 144 槽）与"全天购电费"（Σ 当日各槽结算价 × 购电量）。
- 充放电量：每日 6 个 **真实时钟区间** 4 小时块（题目表2 口径，起点 0:00，沿用模板标签
  0:00-4:00 … 20:00-24:00）。方案丙下第 j 槽 = [10j, 10j+10) 分钟、模型窗口 = [0:10, 24:10)，
  故时钟日 [0:00, 24:00) 的槽集合为
      0:00—4:00 ＝ 前一交易日第 144 槽（＝本日 0:00—0:10）＋ 本日槽 1—23
      4:00—8:00 / 8:00—12:00 / 12:00—16:00 / 16:00—20:00 / 20:00—24:00 ＝ 本日槽 24—47 / 48—71 /
      72—95 / 96—119 / 120—143
  **【2026-09-11 订正】** 原实现按「24 槽自然分组」把槽 1—24 填作 0:00-4:00，与真实时钟区间
  相差 10 分钟（槽 1—24 实际覆盖 0:10—4:10），12 个汇总数中 8 个不符题目要求。
  时刻 0:00 / 24:00 的储电量按字面锚定取值（0:00 = 前一日 E_{d,143}，24:00 = 当日 E_{d,143}）。
  订正后六个时钟块的 SOC 链须自 0:00 精确闭合到 24:00（脚本内逐日断言，见 G9/G10）。
- 紧急购电量：仅列出现紧急购电的日期；连续紧急时段合并为区间（沿用模板表4 格式）。
- 数据来源轨（可执行＝预报驱动口径，D2）：result2←Q2_可执行、result3←Q3_可执行A3、
  result4-2←Q4-2_可执行_价格不可知、result4-3←Q4-3_可执行A3_价格不可知；
  result1 为问题1 单日（2025-01-01）问题1 模型解（mode="q1"，日循环 + 字面 24:00 锚定 6000 kWh）。

用法：
  python "03_代码/正式实验/写模板.py"
"""

from __future__ import annotations

import csv
import datetime
import json
import sys
from pathlib import Path

import numpy as np

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

BASE = Path(r"C:\Users\John\Desktop\26题")
sys.path.insert(0, str(BASE / "03_代码" / "最小运行"))

from openpyxl import load_workbook  # noqa: E402

from 口径 import (ATT, DT, ETA, E_START_1JAN, N_SLOT, ROOT, load_attachment1,  # noqa: E402
                  load_attachment4, sha256)
from 模型 import solve_day  # noqa: E402
from 正式运行 import DEC_END, FEB_START  # noqa: E402

OUT = ROOT / "04_结果与检验" / "results"
DST = OUT                                     # 题目要求的结果文件按《目录与交付物映射.md》第五节放在 results/
TPL = ATT / "附件5"
LOG = ROOT / "04_结果与检验" / "检验日志"

# ---- 题目表2 的真实时钟区间（方案丙）------------------------------------------------
# 第 j 槽 = [10j, 10j+10) 分钟；模型窗口 = [0:10, 24:10)。时钟日 [0:00, 24:00) 恰含 144 槽：
#   块0 = 前一交易日第 144 槽（本日 0:00—0:10）+ 本日槽 1—23；块1—5 = 本日本日槽 24—47 … 120—143。
CLOCK_LABELS = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
CLOCK_TAIL = 144                       # 1-based：本日 0:00—0:10 子槽（属前一交易日）
CLOCK_HEAD = (1, 23)                   # 1-based 闭区间：本日 0:10—4:00
CLOCK_BLOCKS = [None, (24, 47), (48, 71), (72, 95), (96, 119), (120, 143)]  # 1-based 闭区间
FIELD_MAP = {"plan": "plan_kwh", "adj": "adj_kwh", "C": "charge_kwh", "D": "discharge_kwh",
             "e": "emg_kwh", "w": "curtail_kwh", "Eend": "E_end_kwh"}


_buf: list[str] = []
_checks: list[tuple[str, bool, str]] = []


def log(*parts):
    line = " ".join(str(p) for p in parts)
    print(line)
    _buf.append(line)


def record(name: str, ok: bool, detail: str = ""):
    _checks.append((name, bool(ok), detail))
    log(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f" | {detail}" if detail else ""))


# ---------------------------------------------------------------- 读取正式运行产物
def read_slots(track: str) -> dict:
    """逐槽明细 → {date: {plan/adj/C/D/e/w/Eend: ndarray(144)}}"""
    path = OUT / f"正式运行_逐槽_{track}.csv"
    if not path.exists():
        raise FileNotFoundError(f"缺少逐槽文件：{path}")
    data: dict = {}
    with path.open(encoding="utf-8-sig") as fh:
        for r in csv.DictReader(fh):
            d = datetime.date.fromisoformat(r["date"])
            if d not in data:
                data[d] = {k: np.zeros(N_SLOT) for k in FIELD_MAP}
            j = int(r["slot"]) - 1
            for k, col in FIELD_MAP.items():
                data[d][k][j] = float(r[col])
    return data


def read_daily(track: str) -> list[dict]:
    path = OUT / f"正式运行_逐日_{track}.csv"
    if not path.exists():
        raise FileNotFoundError(f"缺少逐日文件：{path}")
    with path.open(encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def day_of(s: str) -> datetime.date:
    return datetime.date.fromisoformat(str(s)[:10])


def content_fingerprint(path) -> str:
    """xlsx 的**内容级**指纹：全部非空单元格（含工作表名、坐标与规范化值）。

    存在理由：实测 openpyxl 写出的 xlsx **字节级不可复现**——同一逻辑连续两次运行，
    5 个文件的 SHA-256 与字节数全部变化（如 result3 799,468→799,467 B），
    但内容级指纹与非空格数完全一致。故「结果数值未被修改」应以本指纹为准，
    字节 SHA-256 仅供参考。
    """
    import hashlib
    wb = load_workbook(path, data_only=True)
    h = hashlib.sha256()
    for ws in wb.worksheets:
        h.update(ws.title.encode("utf-8"))
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                s = f"{c.value:.10g}" if isinstance(c.value, float) else str(c.value)
                h.update(f"{c.coordinate}={s}\x1f".encode("utf-8"))
    wb.close()
    return h.hexdigest().upper()


# ---------------------------------------------------------------- 时钟口径六段
def recover_dawn_slot(e_before: float, e_after: float) -> tuple[float, float]:
    """由两个边界状态反推 [0:00, 0:10) 子槽的 (充电量, 放电量)。

    依据：同时段充放电互斥 + SOC 递推 E_after = E_before + 0.9·C − D/0.9。
    两边界状态已知且 C·D = 0 时，(C, D) 唯一确定。
    用途：交付区间首日（2025-02-01）的 0:00—0:10 属 1 月 31 日第 144 槽，
    而 1 月逐槽明细按约定不持久化；本函数给出可复核的精确恢复值。
    """
    de = e_after - e_before
    return (de / ETA, 0.0) if de >= 0 else (0.0, -de * ETA)


def clock_blocks_of_day(sd: dict, days: list, k: int) -> tuple[list, list]:
    """第 k 日（0-based）六个 **真实时钟区间** 的 (充电量列表, 放电量列表)。"""
    d = days[k]
    C, D = sd[d]["C"], sd[d]["D"]
    head_c = float(C[CLOCK_HEAD[0] - 1:CLOCK_HEAD[1]].sum())
    head_d = float(D[CLOCK_HEAD[0] - 1:CLOCK_HEAD[1]].sum())
    if k > 0:
        Cp, Dp = sd[days[k - 1]]["C"], sd[days[k - 1]]["D"]
        tail_c, tail_d = float(Cp[CLOCK_TAIL - 1]), float(Dp[CLOCK_TAIL - 1])
    else:
        tail_c, tail_d = sd[d]["dawn_C"], sd[d]["dawn_D"]
    out_c, out_d = [tail_c + head_c], [tail_d + head_d]
    for (s, e_) in CLOCK_BLOCKS[1:]:
        out_c.append(float(C[s - 1:e_].sum()))
        out_d.append(float(D[s - 1:e_].sum()))
    return out_c, out_d


def clock_soc_chain(sd: dict, days: list) -> float:
    """逐日六段时钟 SOC 链自 0:00 到 24:00 的最大闭合误差（kWh）。"""
    worst = 0.0
    for k, d in enumerate(days):
        c_, d_ = clock_blocks_of_day(sd, days, k)
        soc = float(sd[d]["E_clock_0"])
        for bi in range(6):
            soc = soc + ETA * c_[bi] - d_[bi] / ETA
        worst = max(worst, abs(soc - float(sd[d]["E_clock_24"])))
    return worst



# ---------------------------------------------------------------- 通用写入
def clear_rows(ws, r0: int, r1: int, ncol: int):
    """清空 [r0, r1] 行的前 ncol 列。

    注意：必须用 `cell(...).value = None` **显式赋值**——openpyxl 的 `cell(..., value=None)`
    与直接读取一样不会改动单元格（P2 质检 P1-1 实测：用 value=None 传参时清空是空操作，
    导致模板示例行残留 '⁝'、'2025-02-02'、'2025-12-31'、'24:00' 等旧内容）。
    """
    for r in range(r0, r1 + 1):
        for c in range(1, ncol + 1):
            ws.cell(row=r, column=c).value = None


def fill_plan_sheet(ws, days, qty: dict, price_of):
    """计划购电量/调整购电量：col1 日期，col2..145 槽值，col146 全天购电量，col147 全天购电费。"""
    for k, d in enumerate(days):
        r = 2 + k
        ws.cell(row=r, column=1, value=datetime.datetime(d.year, d.month, d.day))
        q = qty[d]
        p = price_of(d)
        for j in range(N_SLOT):
            ws.cell(row=r, column=2 + j, value=round(float(q[j]), 4))
        ws.cell(row=r, column=146, value=round(float(np.sum(q)), 4))
        ws.cell(row=r, column=147, value=round(float(np.sum(p * q)), 4))


def fill_cd_sheet(ws, days, sd: dict, date_col: bool):
    """充放电量：每日 6 行（**真实时钟区间** 4 小时块）；时刻 0:00/24:00 行带储电量。
    列偏移：date_col=True 时 1=日期、2=时间段、3=充电量、4=放电量、5=时刻、6=储电量；
    date_col=False（result1）时 1=时间段、2=充电量、3=放电量、4=时刻、5=储电量。
    时间段标签先从模板首个日块读出（保持与官方模板逐字一致），再为**每个日块**写入
    ——否则模板示例行之外的日块会缺失 '0:00-4:00'…'20:00-24:00' 标签（G8 残留扫描发现）。
    【2026-09-11 订正】块内槽集合改为真实时钟区间（见 CLOCK_BLOCKS 与 clock_blocks_of_day）。
    """
    c0 = 1 if date_col else 0
    labels = [ws.cell(row=2 + b, column=c0 + 1).value for b in range(len(CLOCK_LABELS))]
    if any(v is None for v in labels):
        raise RuntimeError("充放电量模板缺少 4 小时段标签，无法按官方格式填写")
    if [str(v) for v in labels] != CLOCK_LABELS:
        raise RuntimeError(f"充放电量模板段标签与题目表2 时钟区间不符：{labels}")
    if not date_col:
        clear_rows(ws, 2, max(ws.max_row, 1 + 6 * len(days)), 5)
    for k, d in enumerate(days):
        base = 2 + 6 * k
        blk_c, blk_d = clock_blocks_of_day(sd, days, k)
        for b in range(6):
            r = base + b
            if date_col:
                # 仅日块首行写日期；**续行必须显式清空**——模板的 12-31 示例块首行会落在
                # 新布局的某日续行上（P2 修正版复核 P1-R2a：result2/4-2 的 R15C1、result3/4-3 的 R21C1
                # 残留 '2025-12-31'），不显式清空会留下合法但错误的日期。
                ws.cell(row=r, column=1,
                        value=datetime.datetime(d.year, d.month, d.day) if b == 0 else None)
                if b:
                    ws.cell(row=r, column=1).value = None
            ws.cell(row=r, column=c0 + 1, value=labels[b])
            ws.cell(row=r, column=c0 + 2, value=round(blk_c[b], 4))
            ws.cell(row=r, column=c0 + 3, value=round(blk_d[b], 4))
            if b >= 2:
                ws.cell(row=r, column=c0 + 4).value = None      # 显式清空（value=None 传参是空操作）
                ws.cell(row=r, column=c0 + 5).value = None
        ws.cell(row=base, column=c0 + 4, value=datetime.time(0, 0))
        ws.cell(row=base, column=c0 + 5, value=round(float(sd[d]["E_clock_0"]), 4))
        ws.cell(row=base + 1, column=c0 + 4, value="24:00")
        ws.cell(row=base + 1, column=c0 + 5, value=round(float(sd[d]["E_clock_24"]), 4))


def emergency_intervals(e_arr: np.ndarray) -> list[tuple[int, int, float]]:
    """连续紧急槽合并为 (起始分钟, 结束分钟, kWh) 区间；槽 j 覆盖 [10(j+1), 10(j+2)) 分钟。"""
    out = []
    t = 0
    while t < N_SLOT:
        if e_arr[t] > 1e-9:
            s = t
            while t + 1 < N_SLOT and e_arr[t + 1] > 1e-9:
                t += 1
            out.append((10 * (s + 1), 10 * (t + 2), float(e_arr[s:t + 1].sum())))
        t += 1
    return out


def fill_emg_sheet(ws, days, sd: dict) -> tuple[int, float]:
    """紧急购电量：仅列有紧急购电的日期；日期只写在当日首行。"""
    clear_rows(ws, 2, max(ws.max_row, 2), 3)
    row = 2
    total = 0.0
    for d in days:
        ivs = emergency_intervals(sd[d]["e"])
        for k, (m0, m1, kwh) in enumerate(ivs):
            if k == 0:
                ws.cell(row=row, column=1, value=datetime.datetime(d.year, d.month, d.day))
            else:
                ws.cell(row=row, column=1).value = None         # 同日后继行必须显式清空
            ws.cell(row=row, column=2, value=f"{m0 // 60}:{m0 % 60:02d}-{m1 // 60}:{m1 % 60:02d}")
            ws.cell(row=row, column=3, value=round(kwh, 4))
            total += kwh
            row += 1
    return row - 2, total


# ---------------------------------------------------------------- result1（问题1 单日）
def write_result1() -> dict:
    a1 = load_attachment1()
    q1 = solve_day(a1["price"], a1["load_kw"], a1["pv_kw"], E_start=None, mode="q1",
                   E_clock_24=E_START_1JAN)
    if not q1["ok"]:
        raise RuntimeError(f"问题1 求解失败：{q1['status']} {q1['message']}")
    wb = load_workbook(TPL / "result1.xlsx")
    ws = wb["计划购电量"]
    for j in range(N_SLOT):
        ws.cell(row=2 + j, column=2, value=round(float(q1["b"][j]), 4))
    ws_cd = wb["充放电量"]
    # 【2026-09-11 订正】真实时钟区间：块0 = 槽144 + 槽1—23，块1—5 = 槽24—47 … 槽120—143。
    # 问题1 为「典型日周期重复」，前一日第 144 槽即本日第 144 槽，与 result1_时钟口径修正版.py 一致。
    C1, D1 = q1["C"], q1["D"]
    blk_c = [float(C1[CLOCK_TAIL - 1]) + float(C1[CLOCK_HEAD[0] - 1:CLOCK_HEAD[1]].sum())]
    blk_d = [float(D1[CLOCK_TAIL - 1]) + float(D1[CLOCK_HEAD[0] - 1:CLOCK_HEAD[1]].sum())]
    for (s, e_) in CLOCK_BLOCKS[1:]:
        blk_c.append(float(C1[s - 1:e_].sum()))
        blk_d.append(float(D1[s - 1:e_].sum()))
    for k in range(6):
        ws_cd.cell(row=2 + k, column=2, value=round(blk_c[k], 4))
        ws_cd.cell(row=2 + k, column=3, value=round(blk_d[k], 4))
    e00 = float(E_START_1JAN)          # 字面 0:00（= 前一日 E_{d,143}，问题1 由锚定给出 6000）
    e24 = float(q1["E_all"][-2])       # 字面 24:00 = 当日 E_{d,143}
    ws_cd.cell(row=2, column=5, value=round(e00, 4))
    ws_cd.cell(row=3, column=5, value=round(e24, 4))
    dst = DST / "result1.xlsx"
    wb.save(dst)
    wb.close()

    # 时钟口径 SOC 链自 0:00 闭合到 24:00（问题1 的独立判据）
    soc = e00
    for k in range(6):
        soc = soc + ETA * blk_c[k] - blk_d[k] / ETA
    chain_err = abs(soc - e24)

    rb = load_workbook(dst, data_only=True)
    p = rb["计划购电量"]
    cd = rb["充放电量"]
    ok_slots = all(p.cell(row=2 + j, column=2).value is not None for j in range(N_SLOT))
    ok_label = p.cell(row=61, column=1).value == "10:00-10:10" and p.cell(row=145, column=1).value == "0:00+1-0:10+1"
    ok_val = abs(float(p.cell(row=61, column=2).value) - q1["b"][59]) < 1e-6
    cd_labels = [str(cd.cell(row=2 + b, column=1).value) for b in range(6)]
    # 容差 1e-4：模板按 4 位小数存储，单值舍入上界为 5e-5，取 2 倍留裕度（实测最大偏差恰为 5.0e-5）。
    ok_cd = all(abs(float(cd.cell(row=2 + b, column=2).value) - blk_c[b]) < 1e-4 for b in range(6)) and \
        all(abs(float(cd.cell(row=2 + b, column=3).value) - blk_d[b]) < 1e-4 for b in range(6)) and \
        abs(float(cd.cell(row=2, column=5).value) - e00) < 1e-6 and \
        abs(float(cd.cell(row=3, column=5).value) - e24) < 1e-6
    record("G1 result1 落盘且回读一致（144 槽 + 6 个时钟区间 + 0:00/24:00 储电量）",
           ok_slots and ok_label and ok_val and ok_cd,
           f"费用 {q1['objective']:.4f} 元；0:00={e00:.4f}、24:00={e24:.4f} kWh；"
           f"E0={q1['E_all'][0]:.4f} E144={q1['E_all'][-1]:.4f}")
    record("G1b result1 六段为真实时钟区间且 SOC 链自 0:00 闭合到 24:00",
           cd_labels == CLOCK_LABELS and chain_err < 1e-6,
           f"标签={cd_labels}；链误差 {chain_err:.3e} kWh")
    rb.close()
    return {"file": dst, "objective": q1["objective"]}


# ---------------------------------------------------------------- result2/4-2 与 result3/4-3
def write_year_template(track: str, tpl_name: str, dst_name: str, price_src: str,
                        with_adjust: bool) -> dict:
    daily = read_daily(track)
    sd = read_slots(track)
    days = [day_of(r["date"]) for r in daily]
    if len(days) != 334:
        record(f"{dst_name} 天数检查", False, f"逐日文件天数 {len(days)}（期望 334）")
    price_days = None
    if price_src == "附件1":
        a1 = load_attachment1()
        price_of = lambda d: a1["price"]  # noqa: E731
    else:
        a4 = load_attachment4()
        price_days = a4
        price_of = lambda d: a4["price"][a4["index"][d]]  # noqa: E731

    for r in daily:
        d = day_of(r["date"])
        sd[d]["E_clock_0"] = float(r["E_clock_0"])
        sd[d]["E_actual_start"] = float(r["E_actual_start"])
        sd[d]["E_clock_24"] = float(r["E_clock_24"])
        sd[d]["plan_ref"] = float(r["plan_kwh"])
        sd[d]["adj_ref"] = float(r["adj_kwh"])
        sd[d]["emg_ref"] = float(r["emg_kwh"])
        sd[d]["cost_plan_ref"] = float(r["cost_plan"])

    # ---- 0:00—0:10 子槽（属前一交易日第 144 槽）--------------------------------------
    # 交付首日 2025-02-01 的子槽属 2025-01-31 第 144 槽，1 月逐槽明细按约定不持久化，
    # 故由两个边界状态（0:00 与 0:10 储电量）精确恢复；其余 333 日用持久化值校验该恢复法。
    worst_rec = 0.0
    for k, d in enumerate(days):
        c_r, d_r = recover_dawn_slot(float(sd[d]["E_clock_0"]), float(sd[d]["E_actual_start"]))
        if k > 0:
            Cp, Dp = sd[days[k - 1]]["C"], sd[days[k - 1]]["D"]
            worst_rec = max(worst_rec, abs(c_r - float(Cp[CLOCK_TAIL - 1])),
                            abs(d_r - float(Dp[CLOCK_TAIL - 1])))
        else:
            sd[d]["dawn_C"], sd[d]["dawn_D"] = c_r, d_r
    record(f"G9 {dst_name} 首日 0:00—0:10 子槽由边界状态恢复；与其余 333 日持久化值一致",
           worst_rec < 1e-6,
           f"333 日最大偏差 {worst_rec:.3e} kWh；首日取 充 {sd[days[0]]['dawn_C']:.4f} / "
           f"放 {sd[days[0]]['dawn_D']:.4f} kWh")
    chain_err = clock_soc_chain(sd, days)
    record(f"G10 {dst_name} 六个时钟区间 SOC 链自 0:00 精确闭合到 24:00（334 日）",
           chain_err < 1e-6, f"最大链误差 {chain_err:.3e} kWh")

    wb = load_workbook(TPL / tpl_name)
    plan_qty = {d: sd[d]["plan"] for d in days}
    fill_plan_sheet(wb["计划购电量"], days, plan_qty, price_of)
    if with_adjust:
        adj_qty = {d: sd[d]["adj"] for d in days}
        fill_plan_sheet(wb["调整购电量"], days, adj_qty, price_of)
    fill_cd_sheet(wb["充放电量"], days, sd, date_col=True)
    n_emg, emg_sum = fill_emg_sheet(wb["紧急购电量"], days, sd)
    dst = DST / dst_name
    wb.save(dst)
    wb.close()

    # ---- 回读核验 ----
    rb = load_workbook(dst, data_only=True)
    wsb = rb["计划购电量"]
    wsa = rb["调整购电量"] if with_adjust else None
    wsc = rb["充放电量"]
    wse = rb["紧急购电量"]
    max_plan = max_adj = 0.0
    max_tot = max_fee = 0.0
    dates_ok = True
    for k, d in enumerate(days):
        r = 2 + k
        v = wsb.cell(row=r, column=1).value
        dates_ok &= (hasattr(v, "date") and v.date() == d)
        arr = np.array([float(wsb.cell(row=r, column=2 + j).value) for j in range(N_SLOT)])
        max_plan = max(max_plan, float(np.abs(arr - sd[d]["plan"]).max()))
        max_tot = max(max_tot, abs(float(wsb.cell(row=r, column=146).value) - float(arr.sum())))
        p = price_of(d)
        max_fee = max(max_fee, abs(float(wsb.cell(row=r, column=147).value) - float(np.sum(p * arr))))
        if wsa is not None:
            arr_a = np.array([float(wsa.cell(row=r, column=2 + j).value) for j in range(N_SLOT)])
            max_adj = max(max_adj, float(np.abs(arr_a - sd[d]["adj"]).max()))
    # 充放电量回读（首日与末日各抽验）：列 3/4=充电量/放电量，列 5=时刻，列 6=储电量
    cd_ok = True
    for k in (0, len(days) - 1):
        d = days[k]
        base = 2 + 6 * k
        blk_c, blk_d = clock_blocks_of_day(sd, days, k)
        for b in range(6):
            cw = float(wsc.cell(row=base + b, column=3).value)
            dw = float(wsc.cell(row=base + b, column=4).value)
            cd_ok &= abs(cw - blk_c[b]) < 1e-3
            cd_ok &= abs(dw - blk_d[b]) < 1e-3
        cd_ok &= abs(float(wsc.cell(row=base, column=6).value) - sd[d]["E_clock_0"]) < 1e-3
        cd_ok &= abs(float(wsc.cell(row=base + 1, column=6).value) - sd[d]["E_clock_24"]) < 1e-3
        cd_ok &= str(wsc.cell(row=base + 1, column=5).value) == "24:00"
    # 紧急购电量回读：逐行合计 vs 逐日文件合计
    emg_file = sum(sd[d]["emg_ref"] for d in days)
    emg_read = 0.0
    r = 2
    while wse.cell(row=r, column=3).value is not None:
        emg_read += float(wse.cell(row=r, column=3).value)
        r += 1
    fee_sum = sum(float(wsb.cell(row=2 + k, column=147).value) for k in range(len(days)))
    cost_plan_sum = sum(sd[d]["cost_plan_ref"] for d in days)
    record(f"G2 {dst_name} 计划购电量 334 行×144 槽回读一致", dates_ok and max_plan < 1e-4 and max_tot < 1e-2,
           f"日期齐全={dates_ok} 最大偏差={max_plan:.2e} 行合计偏差={max_tot:.2e}")
    if with_adjust:
        record(f"G3 {dst_name} 调整购电量回读一致", max_adj < 1e-4, f"最大偏差={max_adj:.2e}")
    record(f"G4 {dst_name} 充放电量回读一致（首末日在 6 个时钟区间与 0:00/24:00 上抽验）", cd_ok)
    # 残留内容扫描（P2 质检 P1-1 的针对性回归）：占位符、越界日期、错位时刻
    residual = []
    for ws in rb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.strip() == "⁝":
                    residual.append(f"{ws.title}!{cell.coordinate} 残留占位符")
    cd_res = rb["充放电量"]
    c0 = 0 if "日期" not in [cd_res.cell(row=1, column=1).value] else 1
    for k in range(len(days)):
        base = 2 + 6 * k
        # 年份模板的充放电量日为 6 行块：日期只在块首行（列 1），续行必须为空
        for b in range(6):
            r = base + b
            v1 = cd_res.cell(row=r, column=1).value
            if b == 0:
                if not hasattr(v1, "date") or v1.date() != days[k]:
                    residual.append(f"充放电量!A{r} 日块首行日期错误 {v1}")
            elif v1 is not None:
                residual.append(f"充放电量!A{r} 续行应空却为 {v1}")
        for b in range(6):
            if cd_res.cell(row=base + b, column=c0 + 1).value in (None, ""):
                residual.append(f"充放电量!R{base + b}C{c0 + 1} 缺时间段标签")
            if cd_res.cell(row=base + b, column=c0 + 2).value is None:
                residual.append(f"充放电量!R{base + b}C{c0 + 2} 缺充电量")
        for b in range(2, 6):
            for col in (c0 + 4, c0 + 5):
                if cd_res.cell(row=base + b, column=col).value is not None:
                    residual.append(f"充放电量!R{base + b}C{col} 应空却非空")
        for col in (c0 + 4, c0 + 5):
            if cd_res.cell(row=base, column=col).value is None:
                residual.append(f"充放电量!R{base}C{col} 应有值却为空")
        if cd_res.cell(row=base + 1, column=c0 + 4).value != "24:00":
            residual.append(f"充放电量!R{base + 1}C{c0 + 4} 24:00 标签缺失")
    emg_ws = rb["紧急购电量"]
    prev_date = None
    r = 2
    while emg_ws.cell(row=r, column=3).value is not None:
        v1 = emg_ws.cell(row=r, column=1).value
        if v1 is not None:
            if not hasattr(v1, "date") or not (FEB_START <= v1.date() <= DEC_END):
                residual.append(f"紧急购电量!A{r} 非法日期 {v1}")
            prev_date = v1
        elif prev_date is None:
            residual.append(f"紧急购电量!A{r} 首行缺日期")
        r += 1
    record(f"G8 {dst_name} 无模板残留（占位符／错位时刻／越界日期）", not residual,
           f"残留项 {len(residual)}" + (f"：{residual[:4]}" if residual else "（占位符、充放电量空行、紧急购电日期分组均正确）"))
    record(f"G5 {dst_name} 紧急购电量区间合计与逐日文件一致",
           abs(emg_read - emg_file) < 1e-3 * max(1, n_emg) and n_emg > 0,
           f"区间行 {n_emg} 行、合计 {emg_read:,.2f} kWh；逐日文件 {emg_file:,.2f} kWh；"
           f"偏差 {abs(emg_read - emg_file):.4f} kWh（4 位小数舍入容差）")
    record(f"G6 {dst_name} 全天购电费合计与逐日文件 cost_plan 一致",
           abs(fee_sum - cost_plan_sum) < 10.0,
           f"表内 {fee_sum:,.2f} 元；逐日文件计划电费合计 {cost_plan_sum:,.2f} 元；"
           f"差 {abs(fee_sum - cost_plan_sum):.2f} 元（逐槽 4 位小数舍入）；逐行最大偏差={max_fee:.2e}")
    if price_days is not None:
        record(f"G7 {dst_name} 价格源为附件4 当日实际电价（价格不可知轨的计划价不用于结算列）", True,
               f"附件4 覆盖 {len(price_days['index'])} 日")
    rb.close()
    return {"file": dst.name, "track": track, "days": len(days), "emg_rows": n_emg, "emg_kwh": emg_read,
            "fee_sum": fee_sum, "cost_plan_file": cost_plan_sum}


def main() -> int:
    DST.mkdir(parents=True, exist_ok=True)
    log("=" * 96)
    log(f"F1 正式模板落盘 | {datetime.datetime.now():%Y-%m-%d %H:%M:%S}")
    log(f"解释器 {sys.executable} | 输出目录 {DST}")
    log(f"输入哈希：附件1 {sha256(ATT / '附件1.xlsx')[:16]} 附件4 {sha256(ATT / '附件4.xlsx')[:16]}")
    log("=" * 96)

    manifest = {}
    r1 = write_result1()
    manifest["result1"] = {"file": r1["file"].name, "track": "问题1 单日 2025-01-01（mode=q1）",
                           "objective_yuan": r1["objective"]}
    manifest["result2"] = write_year_template("Q2_可执行", "result2.xlsx", "result2.xlsx", "附件1", False)
    manifest["result3"] = write_year_template("Q3_可执行A3", "result3.xlsx", "result3.xlsx", "附件1", True)
    manifest["result4-2"] = write_year_template("Q4-2_可执行_价格不可知", "result4-2.xlsx", "result4-2.xlsx",
                                                "附件4", False)
    manifest["result4-3"] = write_year_template("Q4-3_可执行A3_价格不可知", "result4-3.xlsx", "result4-3.xlsx",
                                                "附件4", True)

    log("-" * 96)
    files = {}
    for key in ("result1", "result2", "result3", "result4-2", "result4-3"):
        p = DST / f"{key}.xlsx"
        files[key] = {"bytes": p.stat().st_size, "sha256": sha256(p),
                      "content_sha256": content_fingerprint(p)}
        log(f"  {key}.xlsx  {p.stat().st_size:>9,} B  {sha256(p)[:32]}…  内容指纹 {files[key]['content_sha256'][:32]}…")
    srcs = {}
    for tr in ("Q2_可执行", "Q3_可执行A3", "Q4-2_可执行_价格不可知", "Q4-3_可执行A3_价格不可知"):
        for kind in ("逐日", "逐槽"):
            p = OUT / f"正式运行_{kind}_{tr}.csv"
            if p.exists():
                srcs[p.name] = {"bytes": p.stat().st_size, "sha256": sha256(p)}
    tpl_src = {f"附件5/{n}": sha256(TPL / n) for n in
               ("result1.xlsx", "result2.xlsx", "result3.xlsx", "result4-2.xlsx", "result4-3.xlsx")}
    man = {"阶段": "F1 正式模板落盘", "生成时间": f"{datetime.datetime.now():%Y-%m-%d %H:%M:%S}",
           "来源轨": manifest, "模板输入哈希": tpl_src, "结果文件": files, "数据来源": srcs,
           "指纹说明": "sha256 为文件字节哈希，实测跨次运行不稳定（openpyxl 写出字节不可复现）；"
                       "content_sha256 为全部非空单元格的内容级指纹，跨次运行稳定，"
                       "判定「结果数值是否被修改」应以 content_sha256 为准。",
           "口径": {
               "计划购电量": "原始 0:00 日前计划 b_t（144 槽；末两列=全天购电量/全天购电费）",
               "调整购电量": "相对原始 0:00 计划的调整后购电量 a_t（问题3/4-3 直接调整规则）",
               "全天购电费": "Σ(当日各槽结算价 × 购电量)；问题2/3 用附件1 电价，问题4 用附件4 当日实际电价",
               "充放电量": "每日 6 个真实时钟区间（题目表2）：0:00-4:00 = 前一交易日第144槽 + 本日槽1-23；"
                           "4:00-8:00/8:00-12:00/12:00-16:00/16:00-20:00/20:00-24:00 = 本日槽 24-47/48-71/72-95/96-119/120-143；"
                           "0:00/24:00 为字面锚定储电量，六段 SOC 链自 0:00 精确闭合到 24:00（2026-09-11 订正）",
               "紧急购电量": "仅列出现紧急购电的日期，连续时段合并为区间",
           }}
    mpath = OUT / "复现清单_F1模板.json"
    mpath.write_text(json.dumps(man, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    log(f"  复现清单 {mpath.name}  sha256 {sha256(mpath)[:32]}…")

    n_fail = sum(1 for _, ok, _ in _checks if not ok)
    log("=" * 96)
    log(f"检查合计 {len(_checks)} 项，FAIL {n_fail} 项")
    log("=" * 96)
    LOG.mkdir(parents=True, exist_ok=True)
    (LOG / "正式模板落盘.txt").write_text("\n".join(_buf) + "\n", encoding="utf-8")
    return 0 if n_fail == 0 else 2


if __name__ == "__main__":
    raise SystemExit(main())
