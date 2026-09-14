# -*- coding: utf-8 -*-
"""问题4 口径B（0:00 不知道当日电价）核心模块。

与原口径A 的唯一区别（改造方案 §5.A）：**决策用预测价 `price_dec`，结算用真实价 `price_settle`**。
- 日前计划 LP、6/12/18 更新 LP、10 分钟执行器 causal_step 一律使用 `price_dec`；
- 账单 `pf = p_settle@q0`、`ef = 5*p_settle@qEM`、调整费按 `p_settle` 的增量重算；
- 跨日库存续存价值 v 为**逐日滚动**（L1 规则），用因果的价格预测构造；
- 原实现 `ValuePolicy.value_for` 用 `fp[di+1]`（含当日实际价）属前视，本模块改为 `fp[di]`（见 `残余问题.md`）。

复用的部分（与 q23_janwarm_v4 一致）：1 月因果冷启动、整日配对残差场景、S=10/seed、因果执行器、
Q3 式更新组合与历史选型部署、结算基准（相对上次确认计划的序贯增量）。
"""
from __future__ import annotations

import hashlib
import json
import os
import sys
import time
from pathlib import Path

import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))
from p10_core import (DT, E_INIT, E_MAX, E_MIN, ETA_C, ETA_D, P_MAX_KWH, TP_PRICE,  # noqa: E402
                      causal_step, fc3_step_day, fmt_min, load_all, solve_plan_lp)
from p15_forecasts import baseline_mean7, rolling_forecast, select_past_only  # noqa: E402

PROTOCOL = "q4_scenarioB_v1"
COMBOS = [(), (6,), (12,), (18,), (6, 12), (6, 18), (12, 18), (6, 12, 18)]
ETA = ETA_C
V_OFFPEAK_SLOTS = 30          # 0:00—5:00 共 30 槽（与 L1 规则一致）


def key(combo):
    return "+".join(map(str, combo)) or "none"


def phase_of(di):
    if di < 31:
        return "init"
    if di < 90:
        return "calibration"
    return "eval"


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def code_hashes():
    files = [HERE / n for n in ("p10_core.py", "p15_forecasts.py", "p16_price_forecast.py",
                                "p102_q4_core.py")]
    return {str(p): sha256(p) for p in files if p.exists()}


def input_hashes():
    ad = ROOT / "00_题目与原始数据" / "C题" / "附件"
    return {str(p): sha256(p) for p in sorted(ad.glob("附件[1-4].xlsx"))}


def write_json(path, obj):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=2, default=str)


def save_pickle(path, obj):
    import pickle
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "wb") as f:
        pickle.dump(obj, f, protocol=pickle.HIGHEST_PROTOCOL)


# ---------------------------------------------------------------- 负荷/光伏预测（同 v4 算法）
def make_loadpv_forecasts(out_dir, reuse=True):
    out_dir = Path(out_dir)
    fdir = out_dir / "forecast"
    fdir.mkdir(parents=True, exist_ok=True)
    cache = fdir / "负荷光伏预测缓存.npz"
    meta_path = fdir / "负荷光伏预测缓存_meta.json"
    expected = dict(protocol=PROTOCOL, training_start=14, input_hashes=input_hashes(),
                    code_hashes=code_hashes(), early_policy="lag1_for_di_lt_14", seed=2025)
    if reuse and cache.exists() and meta_path.exists():
        meta = json.loads(meta_path.read_text(encoding="utf-8"))
        if all(meta.get(k) == expected.get(k) for k in expected):
            with np.load(cache) as z:
                if all(k in z for k in ("FL", "FG")):
                    if np.isfinite(z["FL"][:14]).all() and np.isfinite(z["FG"][:14]).all():
                        return str(cache), meta
    data = load_all()
    meta = expected.copy()
    saved, selections, metrics = {}, {}, {}
    for prefix, actual in (("L", data["L"]), ("G", data["G"])):
        lag = np.empty_like(actual)
        lag[0] = actual[0]
        lag[1:] = actual[:-1]
        m7 = baseline_mean7(actual)
        m7[:7] = lag[:7]
        lgb = rolling_forecast(actual, start=14)
        lgb[:14] = lag[:14]
        candidates = {"lgb": lgb, "lag1": lag, "m7": m7}
        forecast, selected = select_past_only(candidates, actual, start=14)
        saved["F" + prefix] = forecast
        selections[prefix] = selected
        for name, arr in candidates.items():
            saved["F" + prefix + "_" + name] = arr
        for name, arr in dict(candidates, online=forecast).items():
            err = arr[31:] - actual[31:]
            metrics[prefix + "_" + name + "_eval"] = dict(
                MAE=float(np.abs(err).mean()), RMSE=float(np.sqrt((err ** 2).mean())),
                bias=float(err.mean()))
    np.savez_compressed(cache, **saved)
    pd.DataFrame(dict(日期=data["dates"], 负荷模型=selections["L"], 光伏模型=selections["G"])).to_csv(
        fdir / "逐日预测选型.csv", index=False, encoding="utf-8-sig")
    meta.update(dict(metrics=metrics, created_at=time.strftime("%Y-%m-%dT%H:%M:%S%z")))
    write_json(meta_path, meta)
    return str(cache), meta


def load_loadpv_forecasts(path):
    with np.load(path) as z:
        f = {k: z[k].copy() for k in ("FL", "FG")}
    if not np.isfinite(f["FL"][:14]).all() or not np.isfinite(f["FG"][:14]).all():
        raise ValueError("1月早期预测不完整")
    return f


# ---------------------------------------------------------------- 价格序列与滚动 v
def price_dec(price_fc: dict, track: str, di: int, t0: int = 0, refresh_mode: str = "none"):
    """决策用价格序列（从槽 t0 起）。refresh_mode='none' → 严格口径B（0:00 一次定全天）；
    'hourly' → 在 6/12/18 刷新节点改用该时刻的预测版本。"""
    base = np.asarray(price_fc[f"FP_{track}"][di], float)
    if refresh_mode == "none" or t0 == 0:
        return base[t0:]
    h = t0 // 6
    k = f"FP_h{h}"
    if k in price_fc:
        return np.asarray(price_fc[k][di], float)[t0:]
    return base[t0:]


def rolling_v(price_fc: dict, track: str, di: int):
    """跨日库存续存价值：v = mean(0:00—5:00 价格预测)/η，逐日滚动。
    因果性：hist7 等预测只由 < di 天的真实价构成；**不使用 fp[di+1]**（原实现前视 1 天）。"""
    base = np.asarray(price_fc[f"FP_{track}"][di], float)
    return float(base[:V_OFFPEAK_SLOTS].mean() / ETA)


# ---------------------------------------------------------------- 场景与冷启动
def scenarios(data, forecasts, di, release_h, S, seed, method, use_fc3=False):
    """整日配对残差场景。**use_fc3=True 时用附件3 已发布的小时预报**分解为 10 分钟光伏预测
    （并与该预报构造残差），与 `p22_backtest`/`p95` 的机制一致——问题4 的 Q3 型轨
    （combo 不为 None）必须使用附件3，否则"更新"通道会退化为仅重抽场景。

    2026-09-12 晚修正：本函数初版遗漏了 `use_fc3` 分支（恒用负荷/光伏点预测），
    导致 Q4 的 8 组合与部署轨不含附件3 信息、且 D4 的 F 轨效果恒为 0（N−F = 0.00）。
    Q4-2 的 A/B/C 轨（combo=None）在两个引擎中本来都不使用附件3，故不受此缺陷影响。
    """
    lf = np.asarray(forecasts["FL"][di], float)
    if use_fc3:
        gf = fc3_step_day(data["fc3"], di, release_h, method=method)[0]
    else:
        gf = np.asarray(forecasts["FG"][di], float)
    hist = np.arange(14, di, dtype=int)
    if len(hist) == 0:
        return lf[None, :], gf[None, :], [], lf, gf, 0
    rng = np.random.default_rng(np.random.SeedSequence([int(seed), int(di), int(release_h)]))
    days = rng.choice(hist, size=int(S), replace=len(hist) < int(S))
    rl = data["L"][days] - forecasts["FL"][days]
    if use_fc3:
        rg = np.stack([data["G"][d] - fc3_step_day(data["fc3"], int(d), release_h, method=method)[0]
                       for d in days])
    else:
        rg = data["G"][days] - forecasts["FG"][days]
    return np.maximum(0.0, lf[None, :] + rl), np.maximum(0.0, gf[None, :] + rg), days.tolist(), lf, gf, len(hist)


def cold_day(data, di, p_settle, E0):
    L, G = data["L"][di], data["G"][di]
    qem = np.maximum(0.0, (L - G) * DT)
    z = np.zeros(144, dtype=float)
    return dict(
        日期=str(data["dates"][di]), di=di, q=z.copy(), q0=z.copy(), qc=z.copy(),
        c=z.copy(), d=z.copy(), E=np.full(144, float(E0)), qEM=qem,
        E0=float(E0), Eend=float(E0), updates=[], scenario_days={"0": []},
        scenario_pool_counts={"0": 0}, planned_emergency_qty=0.0, planned_emergency_fee=0.0,
        计划费=0.0, 调整费=0.0, 紧急费=float(np.sum(5.0 * p_settle * qem)),
        总费用=float(np.sum(5.0 * p_settle * qem)), 阶段=phase_of(di), cold_start=True,
        v_term_used=0.0, price_mae=float(np.abs(p_settle).mean()) * 0.0,
    )


def simulate_day(data, forecasts, price_fc, di, E0, *, track="hist7", refresh_mode="none",
                 v_mode="rolling", price_source="forecast", v_const=None, update_mode="FM",
                 model="B", combo=None, S=10, seed=2025, lam=.3, method="hold", eps_soc=1e-3,
                 tp=TP_PRICE):
    p_settle = np.asarray(data["P_var"][di], float)          # 结算用真实价
    if di == 0:
        return cold_day(data, di, p_settle, E0)
    use_upd = combo is not None
    # use_fc3（是否用附件3 已发布预报建情景池）默认与 use_upd 一致：
    #   combo=None（Q4-2 的 A/B/C 轨）→ 点预报/持久化（与 Q2 同族）
    #   combo=()（Q4-3 的 none 轨）与各事件轨 → 附件3（与 Q3 家族同族）
    # 2026-09-12 晚（独立复核提出"一个布尔表达了两件事"）：新增**默认关闭**的环境变量开关，
    # 仅用于把 none 轨的信息集切回点预报，以量化该约定对交付数字的影响（敏感性实验）。
    # 不设该变量时行为与冻结版**逐位相同**。
    use_fc3 = use_upd
    if os.environ.get("Q4_NONE_NO_FC3") == "1" and combo is not None and len(combo) == 0:
        use_fc3 = False
    ls, gs, days, lf, gf, pool_n = scenarios(data, forecasts, di, 0, S, seed, method,
                                             use_fc3=use_fc3)

    def pdec(t0):
        """决策价：price_source='actual' → 口径A（0:00 已知当日真实曲线，用于隔离电价信息效应）；
        'forecast' → 口径B（只用 0:00 可得的预测价）。"""
        if price_source == "actual":
            return p_settle[t0:]
        return price_dec(price_fc, track, di, t0, refresh_mode)

    p_dec0 = pdec(0)
    v = (float(v_const) if v_const is not None
         else (rolling_v(price_fc, track, di) if v_mode == "rolling" else 0.0))
    if model == "A":
        ls, gs = lf[None, :], gf[None, :]
    plan = solve_plan_lp(p_dec0, ls, gs, E0, lam=lam if model == "C" else 0.0,
                         eps_soc=eps_soc, tp=tp, v_term=v)
    if plan["status"] != 0:
        raise RuntimeError(plan["msg"])
    q0 = plan["q"].copy()
    qc = q0.copy()
    state = float(E0)
    arrays = {k: np.zeros(144, dtype=float) for k in ("c", "d", "E", "qEM")}
    updates, draws, pools = [], {"0": list(days)}, {"0": int(pool_n)}
    expected_qty = float(np.mean(plan["qEM"].sum(axis=1)))
    expected_fee = float(np.mean(np.sum(5.0 * p_settle[None, :] * plan["qEM"], axis=1)))
    adj = 0.0
    for t in range(144):
        hour = t // 6
        if use_upd and t > 0 and t % 6 == 0 and hour in combo:
            ls, gs, days, lf, gf, pool_n = scenarios(data, forecasts, di, hour, S, seed, method,
                                                     use_fc3=True)
            draws[str(hour)] = list(days)
            pools[str(hour)] = int(pool_n)
            if update_mode == "F":
                # D4 式拆分（Q4 版）：只用新发布的附件3 预报刷新日内预测（执行器视界随之更新），
                # **锁定日前购电计划 q0**，不做市场购电调整 ⇒ 调整费为 0、qc 不变。
                updates.append(dict(
                    tau=t, hour=hour, old=qc[t:].copy(), new=qc[t:].copy(), fee=0.0,
                    E0=float(state), scenario_days=list(days), scenario_pool_n=int(pool_n),
                    expected_emergency_qty=0.0, expected_emergency_fee=0.0, forecast_only=True))
            else:
                old = qc[t:].copy()
                p_du = pdec(t)                                         # 决策价（口径A 真实 / 口径B 预测，可刷新）
                update = solve_plan_lp(p_du, ls[:, t:], gs[:, t:], state, q_old=old,
                                       eps_soc=eps_soc, tp=tp, v_term=v)
                if update["status"] != 0:
                    raise RuntimeError(update["msg"])
                qc[t:] = update["q"]
                plus = np.maximum(qc[t:] - old, 0.0)
                minus = np.maximum(old - qc[t:], 0.0)
                fee = float(p_settle[t:] @ (1.5 * plus - .5 * minus))   # 账单用真实价
                adj += fee
                updates.append(dict(
                    tau=t, hour=hour, old=old, new=qc[t:].copy(), fee=fee, E0=float(state),
                    scenario_days=list(days), scenario_pool_n=int(pool_n),
                    expected_emergency_qty=float(np.mean(update["qEM"].sum(axis=1))),
                    expected_emergency_fee=float(np.mean(np.sum(5.0 * p_settle[t:][None, :] * update["qEM"], axis=1))),
                ))
        p_dt = pdec(t)                                                  # 执行器决策价
        step = causal_step(p_dt, qc[t:], state, data["L"][di, t], data["G"][di, t],
                           lf[t:], gf[t:], eps_soc=eps_soc, tp=tp, v_term=v)
        for k in arrays:
            arrays[k][t] = step[k]
        state = float(step["E"])
    pf = float(p_settle @ q0)
    ef = float(5.0 * p_settle @ arrays["qEM"])
    return dict(
        日期=str(data["dates"][di]), di=di, q=q0, q0=q0, qc=qc, **arrays,
        E0=float(E0), Eend=float(state), updates=updates,
        scenario_days=draws, scenario_pool_counts=pools,
        planned_emergency_qty=expected_qty, planned_emergency_fee=expected_fee,
        计划费=pf, 调整费=adj, 紧急费=ef, 总费用=pf + adj + ef,
        阶段=phase_of(di), cold_start=False, v_term_used=float(v),
        price_mae=float(np.abs(p_dec0 - p_settle).mean()),
    )


# ---------------------------------------------------------------- 汇总与落盘
def summary(recs, start=None, stop=None):
    rows = recs[slice(start, stop)] if start is not None or stop is not None else recs
    return dict(
        天数=len(rows), 计划购电费=float(sum(r["计划费"] for r in rows)),
        调整费=float(sum(r["调整费"] for r in rows)),
        紧急购电费=float(sum(r["紧急费"] for r in rows)),
        紧急购电量=float(sum(np.sum(r["qEM"]) for r in rows)),
        总费用=float(sum(r["总费用"] for r in rows)),
        初SOC=float(rows[0]["E0"]) if rows else None,
        末SOC=float(rows[-1]["Eend"]) if rows else None,
        均值日度价格预测MAE=float(np.mean([r["price_mae"] for r in rows])) if rows else None,
        均值v_term=float(np.mean([r["v_term_used"] for r in rows])) if rows else None,
    )


def daily_frame(recs, label):
    return pd.DataFrame([
        dict(日期=r["日期"], 模型=label, 计划购电费=r["计划费"], 调整费=r["调整费"],
             紧急购电费=r["紧急费"], 紧急购电量=float(np.sum(r["qEM"])), 总费用=r["总费用"],
             末SOC=r["Eend"], 阶段=r["阶段"], v_term=r["v_term_used"], 价格预测MAE=r["price_mae"],
             计划紧急购电量=r["planned_emergency_qty"], 计划紧急购电费=r["planned_emergency_fee"])
        for r in recs])


def deploy_from_shadow(data, forecasts, price_fc, shadows, cfg):
    ordered = [key(c) for c in COMBOS]
    n = len(next(iter(shadows.values())))
    cumulative = {k: 0.0 for k in ordered}
    state = E_INIT
    deployed = []
    for i in range(n):
        di = shadows[ordered[0]][i]["di"]
        selected = "none" if di == 0 else min(ordered, key=lambda k: cumulative[k])
        combo = COMBOS[ordered.index(selected)]
        rec = simulate_day(data, forecasts, price_fc, di, state, combo=combo, **cfg)
        rec["selected_combo"] = selected
        rec["selection_history_end"] = di - 1
        deployed.append(rec)
        state = rec["Eend"]
        for k in ordered:
            cumulative[k] += shadows[k][i]["总费用"]
    return deployed, cumulative


def write_result_workbooks(out_dir, q42_b, q43_deployed, dates_from="2025-02-01"):
    import openpyxl
    tpl = ROOT / "00_题目与原始数据" / "C题" / "附件" / "附件5"
    out_dir = Path(out_dir)
    dates = [d.strftime("%Y/%m/%d") for d in pd.date_range(dates_from, "2025-12-31")]
    buckets = [(0, 24), (24, 48), (48, 72), (72, 96), (96, 120), (120, 144)]
    labels = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]

    def fill_plan(ws, recs, field, include_adj=False):
        for j in range(144):
            ws.cell(1, 2 + j, f"{fmt_min(j * 10)}-{fmt_min((j + 1) * 10)}")
        for i, (dt, r) in enumerate(zip(dates, recs)):
            q = np.asarray(r[field])
            ws.cell(2 + i, 1, dt)
            for j in range(144):
                ws.cell(2 + i, 2 + j, round(float(q[j]), 2))
            ws.cell(2 + i, 146, round(float(q.sum()), 2))
            fee = r["计划费"] + r["调整费"] if include_adj else r["计划费"]
            ws.cell(2 + i, 147, round(float(fee), 2))

    def fill_storage(ws, recs):
        for row in range(2, ws.max_row + 1):
            for col in range(1, 7):
                ws.cell(row, col).value = None
        row = 2
        for dt, r in zip(dates, recs):
            for bi, (a, b) in enumerate(buckets):
                ws.cell(row, 1, dt if bi == 0 else None)
                ws.cell(row, 2, labels[bi])
                ws.cell(row, 3, round(float(np.sum(r["c"][a:b])), 2))
                ws.cell(row, 4, round(float(np.sum(r["d"][a:b])), 2))
                if bi == 0:
                    ws.cell(row, 5, "0:00"); ws.cell(row, 6, round(float(r["E0"]), 2))
                if bi == 1:
                    ws.cell(row, 5, "24:00"); ws.cell(row, 6, round(float(r["Eend"]), 2))
                row += 1

    def fill_emergency(ws, recs):
        for row in range(2, ws.max_row + 1):
            for col in range(1, 4):
                ws.cell(row, col).value = None
        row = 2
        for dt, r in zip(dates, recs):
            qem = np.asarray(r["qEM"]); t = 0; intervals = []
            while t < 144:
                if qem[t] <= 1e-6:
                    t += 1; continue
                a = t; amount = 0.0
                while t < 144 and qem[t] > 1e-6:
                    amount += float(qem[t]); t += 1
                intervals.append((a, t, amount))
            if not intervals:
                ws.cell(row, 1, dt); row += 1
            else:
                for a, b, amount in intervals:
                    ws.cell(row, 1, dt); ws.cell(row, 2, f"{fmt_min(a * 10)}-{fmt_min(b * 10)}")
                    ws.cell(row, 3, round(amount, 2)); row += 1

    wb = openpyxl.load_workbook(tpl / "result4-2.xlsx")
    fill_plan(wb["计划购电量"], q42_b[31:], "q0")
    fill_storage(wb["充放电量"], q42_b[31:])
    fill_emergency(wb["紧急购电量"], q42_b[31:])
    wb.save(out_dir / "result4-2.xlsx")

    wb = openpyxl.load_workbook(tpl / "result4-3.xlsx")
    fill_plan(wb["计划购电量"], q43_deployed[31:], "q0")
    fill_plan(wb["调整购电量"], q43_deployed[31:], "qc", include_adj=True)
    fill_storage(wb["充放电量"], q43_deployed[31:])
    fill_emergency(wb["紧急购电量"], q43_deployed[31:])
    wb.save(out_dir / "result4-3.xlsx")
